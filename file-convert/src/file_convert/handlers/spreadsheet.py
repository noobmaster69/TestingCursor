from __future__ import annotations

import csv
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from file_convert.errors import ConversionError, ValidationError
from file_convert.handlers.base import BaseHandler
from file_convert.models import ConversionJob


class SpreadsheetHandler(BaseHandler):
    @property
    def pairs(self) -> frozenset[tuple[str, str]]:
        return frozenset({("xlsx", "csv"), ("csv", "xlsx")})

    def _run(self, job: ConversionJob, output_path: Path) -> str | None:
        if job.source_format == "xlsx" and job.target_format == "csv":
            return self._xlsx_to_csv(job, output_path)
        if job.source_format == "csv" and job.target_format == "xlsx":
            return self._csv_to_xlsx(job, output_path)
        raise ConversionError(f"Unsupported pair: {job.source_format}→{job.target_format}")

    def _resolve_sheet_name(self, job: ConversionJob, wb) -> str:
        sheet_opt = job.options.get("sheet")
        names = wb.sheetnames
        if sheet_opt is None:
            return names[0]
        if isinstance(sheet_opt, int) or (
            isinstance(sheet_opt, str) and sheet_opt.isdigit()
        ):
            idx = int(sheet_opt)
            if idx < 0 or idx >= len(names):
                raise ValidationError(
                    f"Sheet index {idx} out of range. Available: {', '.join(names)}"
                )
            return names[idx]
        if sheet_opt not in names:
            raise ValidationError(
                f"Sheet '{sheet_opt}' not found. Available: {', '.join(names)}"
            )
        return str(sheet_opt)

    def _xlsx_to_csv(self, job: ConversionJob, output_path: Path) -> str:
        encoding = job.options.get("encoding", "utf-8")
        delimiter = job.options.get("delimiter", ",")

        wb = load_workbook(job.source, read_only=True, data_only=True)
        try:
            sheet_name = self._resolve_sheet_name(job, wb)
            ws = wb[sheet_name]
            with output_path.open("w", encoding=encoding, newline="") as f:
                writer = csv.writer(f, delimiter=delimiter)
                for row in ws.iter_rows(values_only=True):
                    writer.writerow(
                        ["" if cell is None else cell for cell in row],
                    )
        finally:
            wb.close()
        return f"Exported sheet to CSV"

    def _csv_to_xlsx(self, job: ConversionJob, output_path: Path) -> str:
        encoding = job.options.get("encoding", "utf-8")
        delimiter = job.options.get("delimiter", ",")
        sheet_name = job.options.get("sheet") or "Sheet1"

        df = pd.read_csv(job.source, encoding=encoding, delimiter=delimiter)
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name=str(sheet_name)[:31], index=False)
        return f"Wrote {len(df)} rows to XLSX"
