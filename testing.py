
from __future__ import annotations

import argparse
import html
import os
import random
import re
import shutil
import stat
import subprocess
import tempfile
import threading
import time
import queue
import sys
from collections import defaultdict
from datetime import datetime
from contextlib import suppress
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterable, Optional
from urllib.parse import parse_qs, quote_plus, unquote, urljoin, urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
except Exception:
    tk = None
    filedialog = messagebox = ttk = None

import requests
from bs4 import BeautifulSoup
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

try:
    from dotenv import load_dotenv
except ImportError:
    load_dotenv = None

try:
    from lxml import etree
except ImportError:
    etree = None


# ---------------------------------------------------------------------------
# Load .env
# ---------------------------------------------------------------------------

def _load_env():
    env_path = Path(__file__).resolve().parent / ".env"
    if load_dotenv is not None:
        load_dotenv(env_path)
    elif env_path.exists():
        for line in env_path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            if "=" in line:
                key, _, value = line.partition("=")
                os.environ.setdefault(key.strip(), value.strip())


_load_env()

FB_EMAIL = os.environ.get("FB_EMAIL", "")
FB_PASSWORD = os.environ.get("FB_PASSWORD", "")
FACEBOOK_USER_DATA_DIR = os.environ.get("FACEBOOK_USER_DATA_DIR", "")
FACEBOOK_PROFILE_DIRECTORY = os.environ.get("FACEBOOK_PROFILE_DIRECTORY", "")
CHROME_BINARY = os.environ.get("CHROME_BINARY", "")


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

DEFAULT_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/131.0.0.0 Safari/537.36"
)

SOCIAL_PLATFORM_DOMAINS = {
    "facebook.com", "www.facebook.com", "m.facebook.com",
    "mbasic.facebook.com", "l.facebook.com", "lm.facebook.com",
    "messenger.com", "www.messenger.com",
    "instagram.com", "www.instagram.com",
    "linkedin.com", "www.linkedin.com",
    "x.com", "www.x.com", "twitter.com", "www.twitter.com",
    "tiktok.com", "www.tiktok.com",
    "youtube.com", "www.youtube.com",
}

FACEBOOK_REJECT_TOKENS = {
    "/posts/", "/photos/", "/videos/", "/events/", "/groups/",
    "/watch/", "/reel/", "/reels/", "/permalink/", "/permalink.php",
    "/story.php", "/marketplace/", "/share/", "/login", "/plugins/",
    "/ads/library/", "/hashtag/", "/sharer/", "/dialog/",
    "/policies/", "/privacy/", "/help/", "/gaming/", "/business/help/",
}

FACEBOOK_SUBPAGE_SEGMENTS = {
    "about", "about_contact_and_basic_info", "about_profile_transparency",
    "community", "events", "followers", "friends_likes", "groups", "likes",
    "locations", "mentions", "notes", "photos", "posts", "reviews",
    "services", "shop", "timeline", "videos", "visitor_posts",
}

# The two last-resort absolute XPaths you supplied.
FACEBOOK_PHONE_XPATH = (
    "/html/body/div[1]/div/div[1]/div/div[3]/div/div/div[1]/div[1]/div/div/div[4]/div[2]/"
    "div/div[1]/div[2]/div/div/div[1]/div/div[3]/div/div/div/div[2]/div[1]"
)
FACEBOOK_EMAIL_XPATH = (
    "/html/body/div[1]/div/div[1]/div/div[3]/div/div/div[1]/div[1]/div/div/div[4]/div[2]/"
    "div/div[1]/div[2]/div/div/div[1]/div/div[3]/div/div/div/div[2]/div[2]/div[1]/div[2]/"
    "span/span/a"
)
FACEBOOK_SEARCH_BAR_XPATH = (
    "/html/body/div[1]/div/div[1]/div/div[2]/div[3]/div/div/div/div/div/label"
)

FACEBOOK_DETAIL_WAIT_SECONDS = 2.0
FACEBOOK_SEARCH_WAIT_SECONDS = 30
FACEBOOK_MANUAL_LOGIN_WAIT_SECONDS = 120
FACEBOOK_SEARCH_SCROLL_ROUNDS = 6

DEFAULT_LOCATIONS: list[tuple[str, str]] = [
    ("Rockville", "MD"),
    ("Gaithersburg", "MD"),
    ("Arlington", "VA"),
    ("Washington", "DC"),
    ("Reston", "VA"),
    ("Herndon", "VA"),
    ("Alexandria", "VA"),
    ("Baltimore", "MD"),
    ("Hyattsville", "MD"),
    ("Bethesda", "MD"),
]
DEFAULT_RESULTS_PER_CITY = 10
DEFAULT_CANDIDATE_SEARCH_DEPTH = 50

_FB_PHONE_LABELS = {"phone", "mobile", "telephone", "tel", "call"}
_FB_EMAIL_LABELS = {"email", "e-mail"}
_FB_WEBSITE_LABELS = {"website", "web site", "site", "visit website"}
_FB_ADDRESS_LABELS = {"address", "location"}
_FB_ANCESTOR_LEVELS = (2, 3, 4, 5)

BLOCKED_EMAIL_DOMAINS = {
    "example.com", "sentry.io", "wixpress.com", "googleapis.com",
    "google.com", "gstatic.com", "schema.org", "w3.org",
    "yellowpages.com", "yelp.com", "bbb.org", "manta.com",
}

_BLOCKED_WEBSITE_HOSTS = {
    "bing.com", "api.whatsapp.com", "whatsapp.com",
    "maps.google.com", "maps.apple.com", "fbcdn.net",
    "cdninstagram.com", "gstatic.com", "schema.org", "w3.org",
}

EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}")
PHONE_RE = re.compile(
    r"""
    (?<!\d)
    (?:\+?1[\s\-.\(\)]*)?
    \(?
    (?P<area>\d{3})
    \)?
    [\s\-.\(\)]*
    (?P<exchange>\d{3})
    [\s\-.\(\)]*
    (?P<subscriber>\d{4})
    (?!\d)
    """,
    re.VERBOSE,
)


# ---------------------------------------------------------------------------
# Data class
# ---------------------------------------------------------------------------

@dataclass
class Lead:
    company_name: str = ""
    phone: str = ""
    email: str = ""
    address: str = ""
    city: str = ""
    state: str = ""
    website: str = ""
    facebook_page: str = ""
    source: str = ""
    search_keyword: str = ""
    notes: str = ""


# ---------------------------------------------------------------------------
# Logger
# ---------------------------------------------------------------------------

_LOG_SUBSCRIBERS: list[Callable[[str], None]] = []


def add_log_subscriber(callback: Callable[[str], None]):
    if callback not in _LOG_SUBSCRIBERS:
        _LOG_SUBSCRIBERS.append(callback)


def remove_log_subscriber(callback: Callable[[str], None]):
    with suppress(ValueError):
        _LOG_SUBSCRIBERS.remove(callback)


def log(message: str):
    ts = time.strftime("%H:%M:%S")
    line = f"[{ts}]  {message}"
    print(line)
    for callback in list(_LOG_SUBSCRIBERS):
        with suppress(Exception):
            callback(line)


# ---------------------------------------------------------------------------
# Generic helpers
# ---------------------------------------------------------------------------

def normalize_name(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (value or "").lower())


def normalize_phone(raw: str) -> str:
    digits = re.sub(r"\D", "", raw or "")
    if len(digits) == 11 and digits.startswith("1"):
        digits = digits[1:]
    if len(digits) == 10:
        return f"({digits[:3]}) {digits[3:6]}-{digits[6:]}"
    return ""


def extract_emails(text: str) -> list[str]:
    found: list[str] = []
    seen: set[str] = set()
    for email in EMAIL_RE.findall(text or ""):
        cleaned = email.strip(".,;:()[]{}<>").lower()
        domain = cleaned.split("@")[-1]
        if domain in BLOCKED_EMAIL_DOMAINS:
            continue
        if cleaned not in seen:
            seen.add(cleaned)
            found.append(cleaned)
    return found


def extract_phones(text: str) -> list[str]:
    found: list[str] = []
    seen: set[str] = set()
    for match in PHONE_RE.finditer(text or ""):
        phone = normalize_phone(match.group(0))
        if phone and phone not in seen:
            found.append(phone)
            seen.add(phone)
    return found


def make_soup(html_text: str) -> BeautifulSoup:
    with suppress(Exception):
        return BeautifulSoup(html_text, "lxml")
    return BeautifulSoup(html_text, "html.parser")


def build_session() -> requests.Session:
    session = requests.Session()
    retry_kwargs = dict(
        total=3, connect=3, read=3, backoff_factor=0.8,
        status_forcelist=(429, 500, 502, 503, 504),
        raise_on_status=False,
    )
    try:
        retry = Retry(allowed_methods=frozenset(["GET", "HEAD"]), **retry_kwargs)
    except TypeError:
        retry = Retry(method_whitelist=frozenset(["GET", "HEAD"]), **retry_kwargs)
    adapter = HTTPAdapter(max_retries=retry, pool_connections=10, pool_maxsize=10)
    session.mount("http://", adapter)
    session.mount("https://", adapter)
    session.headers.update({
        "User-Agent": DEFAULT_USER_AGENT,
        "Accept-Language": "en-US,en;q=0.9",
        "Cache-Control": "no-cache",
    })
    return session


def fetch_html(session: requests.Session, url: str, timeout: int = 20) -> Optional[str]:
    try:
        response = session.get(url, timeout=timeout)
        if response.status_code >= 400:
            return None
        return response.text
    except requests.RequestException:
        return None


def clean_url(url: str) -> str:
    value = html.unescape((url or "").strip())
    if not value:
        return ""
    if value.startswith("//"):
        value = "https:" + value
    if value.startswith("/"):
        return ""
    parsed = urlparse(value)
    if not parsed.scheme:
        value = "https://" + value
        parsed = urlparse(value)
    if not parsed.netloc:
        return ""
    return value


def decode_tracking_url(url: str) -> str:
    href = html.unescape((url or "").strip())
    if not href:
        return ""
    if href.startswith("/l/?") and "uddg=" in href:
        query = parse_qs(urlparse(href).query)
        target = query.get("uddg", [""])[0]
        if target:
            return html.unescape(unquote(target))
    parsed = urlparse(href)
    host = parsed.netloc.lower()
    if "google" in host and "/url" in parsed.path:
        query = parse_qs(parsed.query)
        for key in ("q", "url"):
            target = query.get(key, [""])[0]
            if target:
                return html.unescape(unquote(target))
    if "facebook.com" in host and parsed.path == "/l.php":
        query = parse_qs(parsed.query)
        target = query.get("u", [""])[0]
        if target:
            return html.unescape(unquote(target))
    return href


def first_non_empty(*values) -> str:
    for value in values:
        text = str(value).strip() if value is not None else ""
        if text and text.lower() != "nan":
            return text
    return ""


def domain_key(url: str) -> str:
    parsed = urlparse(clean_url(url))
    host = parsed.netloc.lower().strip()
    if host.startswith("www."):
        host = host[4:]
    return host


def is_social_platform_url(url: str) -> bool:
    host = domain_key(url)
    normalized_domains = {d[4:] if d.startswith("www.") else d for d in SOCIAL_PLATFORM_DOMAINS}
    return host in normalized_domains


def _is_blocked_website(url: str) -> bool:
    if not url:
        return True
    lower = url.lower()
    if "facebook.com" in lower:
        return True
    if is_social_platform_url(url):
        return True
    host = domain_key(url)
    return any(host == b or host.endswith("." + b) for b in _BLOCKED_WEBSITE_HOSTS)


def _normalize_space(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def dedupe_location_pairs(locations: Iterable[tuple[str, str]]) -> list[tuple[str, str]]:
    unique: list[tuple[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for city, state in locations:
        cleaned_city = _normalize_space(city)
        cleaned_state = _normalize_space(state).upper()
        if not cleaned_city or not cleaned_state:
            continue
        key = (cleaned_city.lower(), cleaned_state.lower())
        if key in seen:
            continue
        unique.append((cleaned_city, cleaned_state))
        seen.add(key)
    return unique


def parse_location_value(raw: str) -> tuple[str, str] | None:
    text = _normalize_space(raw)
    if not text:
        return None

    if "|" in text:
        city_part, state_part = text.rsplit("|", 1)
    elif "," in text:
        city_part, state_part = text.rsplit(",", 1)
    else:
        return None

    city = _normalize_space(city_part.strip(" -"))
    state = _normalize_space(state_part.strip()).upper()
    if not city or not state:
        return None
    return city, state


def parse_locations_text(raw: str) -> list[tuple[str, str]]:
    locations: list[tuple[str, str]] = []
    for chunk in re.split(r"[;\n]+", raw or ""):
        parsed = parse_location_value(chunk)
        if parsed:
            locations.append(parsed)
    return dedupe_location_pairs(locations)


def build_locations_from_args(args) -> list[tuple[str, str]]:
    locations: list[tuple[str, str]] = []

    for raw in getattr(args, "location", []) or []:
        parsed = parse_location_value(raw)
        if parsed:
            locations.append(parsed)

    locations.extend(parse_locations_text(getattr(args, "locations", "") or ""))

    single_city = _normalize_space(getattr(args, "city", "") or "")
    single_state = _normalize_space(getattr(args, "state", "") or "").upper()
    if single_city and single_state:
        locations.append((single_city, single_state))

    if not locations:
        return list(DEFAULT_LOCATIONS)
    return dedupe_location_pairs(locations)


DEFAULT_KEYWORDS: list[str] = ["construction companies"]


def dedupe_keywords(keywords: Iterable[str]) -> list[str]:
    cleaned_keywords: list[str] = []
    seen: set[str] = set()
    for keyword in keywords:
        cleaned = _normalize_space(keyword)
        if not cleaned:
            continue
        key = cleaned.lower()
        if key in seen:
            continue
        cleaned_keywords.append(cleaned)
        seen.add(key)
    return cleaned_keywords


def parse_keywords_text(raw: str) -> list[str]:
    keywords: list[str] = []
    for chunk in re.split(r"[;\n]+", raw or ""):
        normalized = _normalize_space(chunk)
        if normalized:
            keywords.append(normalized)
    return dedupe_keywords(keywords)


def normalize_search_terms(search_terms: str | Iterable[str]) -> list[str]:
    if isinstance(search_terms, str):
        return dedupe_keywords([search_terms]) or list(DEFAULT_KEYWORDS)
    return dedupe_keywords(search_terms) or list(DEFAULT_KEYWORDS)


def deduplicate_leads(leads: Iterable[Lead]) -> list[Lead]:
    unique: list[Lead] = []
    lookup: dict[tuple[str, str, str, str], Lead] = {}

    def identity(lead: Lead) -> tuple[str, str, str, str]:
        city = _normalize_space(lead.city).lower()
        state = _normalize_space(lead.state).lower()
        page = normalize_facebook_page_url(lead.facebook_page or "")
        if page:
            return (city, state, "page", page.lower())
        if lead.phone:
            return (city, state, "phone", normalize_phone(lead.phone))
        if lead.email:
            return (city, state, "email", lead.email.lower())
        return (city, state, "name", normalize_name(lead.company_name))

    def merge(existing: Lead, incoming: Lead):
        if not existing.company_name and incoming.company_name:
            existing.company_name = incoming.company_name
        if not existing.phone and incoming.phone:
            existing.phone = incoming.phone
        if not existing.email and incoming.email:
            existing.email = incoming.email
        if not existing.address and incoming.address:
            existing.address = incoming.address
        if not existing.website and incoming.website:
            existing.website = incoming.website
        if not existing.facebook_page and incoming.facebook_page:
            existing.facebook_page = incoming.facebook_page
        if incoming.source:
            existing.source = ", ".join(_dedupe_ci([existing.source, incoming.source]))
        if incoming.search_keyword:
            parts = [part.strip() for part in (existing.search_keyword or "").split("|") if part.strip()]
            parts.append(incoming.search_keyword)
            existing.search_keyword = " | ".join(dedupe_keywords(parts))
        note_parts = [part.strip() for part in (existing.notes or "").split(" | ") if part.strip()]
        note_parts.extend([part.strip() for part in (incoming.notes or "").split(" | ") if part.strip()])
        existing.notes = " | ".join(_dedupe_ci(note_parts))

    for lead in leads:
        key = identity(lead)
        existing = lookup.get(key)
        if existing is None:
            clone = Lead(**lead.__dict__)
            lookup[key] = clone
            unique.append(clone)
        else:
            merge(existing, lead)

    return unique


def build_location_queries(search_term: str, city: str, state: str) -> list[str]:
    search_term = _normalize_space(search_term)
    city = _normalize_space(city)
    state = _normalize_space(state).upper()
    city_state = f"{city} {state}".strip()

    base_queries = [
        f"{search_term} {city_state}",
        f"{search_term} in {city_state}",
        f"{search_term} near {city_state}",
        f"{search_term} {city}",
        f"{search_term} contractors {city_state}",
        f"{search_term} company {city_state}",
    ]

    cleaned_queries: list[str] = []
    seen: set[str] = set()
    for query in base_queries:
        normalized = _normalize_space(query)
        if not normalized:
            continue
        key = normalized.lower()
        if key in seen:
            continue
        cleaned_queries.append(normalized)
        seen.add(key)
    return cleaned_queries


def _dedupe_ci(values: Iterable[str]) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for value in values:
        cleaned = _normalize_space(str(value))
        if not cleaned:
            continue
        key = cleaned.lower()
        if key not in seen:
            out.append(cleaned)
            seen.add(key)
    return out


def _merge_phone_lists(*collections: Iterable[str]) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for collection in collections:
        for value in collection or []:
            normalized = normalize_phone(value)
            if normalized and normalized not in seen:
                out.append(normalized)
                seen.add(normalized)
    return out


def _merge_email_lists(*collections: Iterable[str]) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for collection in collections:
        for value in collection or []:
            cleaned = (value or "").strip().lower()
            if cleaned and cleaned not in seen:
                out.append(cleaned)
                seen.add(cleaned)
    return out


def safe_driver_get(driver, url: str) -> bool:
    try:
        driver.get(url)
        return True
    except Exception:
        return False


def _save_debug_html(debug_dir: str, name: str, html_text: str):
    if not debug_dir:
        return
    Path(debug_dir).mkdir(parents=True, exist_ok=True)
    stamp = time.strftime("%Y%m%d_%H%M%S")
    safe_name = re.sub(r"[^a-zA-Z0-9_.-]+", "_", name)[:80]
    path = Path(debug_dir) / f"{stamp}_{safe_name}.html"
    with suppress(Exception):
        path.write_text(html_text or "", encoding="utf-8")


def _save_debug_png(driver, debug_dir: str, name: str):
    if not debug_dir:
        return
    Path(debug_dir).mkdir(parents=True, exist_ok=True)
    stamp = time.strftime("%Y%m%d_%H%M%S")
    safe_name = re.sub(r"[^a-zA-Z0-9_.-]+", "_", name)[:80]
    path = Path(debug_dir) / f"{stamp}_{safe_name}.png"
    with suppress(Exception):
        driver.save_screenshot(str(path))


# ---------------------------------------------------------------------------
# Facebook URL helpers
# ---------------------------------------------------------------------------

def normalize_facebook_title(text: str) -> str:
    title = (text or "").strip()
    if not title:
        return ""
    title = html.unescape(title)
    title = re.sub(r"^\(\d+\)\s*", "", title)
    title = re.sub(r"\s*[|\-\u2013]\s*Facebook.*$", "", title, flags=re.IGNORECASE)
    title = re.sub(r"\s*[|\-\u2013]\s*Meta.*$", "", title, flags=re.IGNORECASE)
    title = re.sub(r"^Facebook -? ", "", title, flags=re.IGNORECASE)
    return title.strip(" -|\u2013")


def normalize_facebook_page_url(url: str) -> str:
    href = decode_tracking_url(url)
    href = clean_url(href)
    if not href:
        return ""
    parsed = urlparse(href)
    host = parsed.netloc.lower()
    if "facebook.com" not in host:
        return ""
    path = parsed.path or ""
    lower_path = path.lower()
    lower_full = lower_path + ("?" + parsed.query.lower() if parsed.query else "")
    if any(token in lower_full for token in FACEBOOK_REJECT_TOKENS):
        return ""
    if lower_path.rstrip("/") in {
        "", "/home.php", "/pages", "/pg", "/search", "/directory",
        "/friends", "/messages", "/notifications", "/bookmarks",
        "/saved", "/feeds", "/gaming", "/weather", "/settings",
        "/memories", "/fundraisers", "/ads", "/reels",
        "/stories", "/stories/create", "/onthisday", "/policies",
        "/business", "/recover", "/checkpoint", "/composer",
    }:
        return ""
    if lower_path == "/profile.php":
        profile_id = parse_qs(parsed.query).get("id", [""])[0]
        if not profile_id:
            return ""
        return f"https://www.facebook.com/profile.php?id={profile_id}"

    segments = [s for s in path.split("/") if s]
    if not segments:
        return ""
    if segments[0].lower() in {
        "groups", "events", "people", "watch", "search", "hashtag", "marketplace",
        "friends", "messages", "notifications", "bookmarks", "saved", "feeds",
        "gaming", "weather", "settings", "memories", "fundraisers", "ads", "reels",
    }:
        return ""
    while segments and segments[-1].lower() in FACEBOOK_SUBPAGE_SEGMENTS:
        segments.pop()
    if not segments:
        return ""
    cleaned_path = "/" + "/".join(segments)
    return f"https://www.facebook.com{cleaned_path}"


def facebook_slug_to_label(url: str) -> str:
    normalized = normalize_facebook_page_url(url)
    if not normalized:
        return ""
    parsed = urlparse(normalized)
    segments = [s for s in parsed.path.split("/") if s]
    if not segments:
        return ""
    if segments[0].lower() == "profile.php":
        return ""
    candidate = segments[-1]
    if candidate.isdigit() and len(segments) > 1:
        candidate = segments[-2]
    if segments[0].lower() == "pg" and len(segments) > 1:
        candidate = segments[1]
    elif segments[0].lower() == "pages" and len(segments) > 1:
        if len(segments) >= 3 and segments[-1].isdigit():
            candidate = segments[-2]
        else:
            candidate = segments[1]
    candidate = re.sub(r"[-_]+", " ", candidate)
    candidate = re.sub(r"\s+", " ", candidate).strip()
    return candidate


def build_facebook_page_variants(page_url: str) -> list[str]:
    normalized = normalize_facebook_page_url(page_url)
    if not normalized:
        return []
    parsed = urlparse(normalized)
    host = "www.facebook.com"

    variants: list[str] = []
    if parsed.path == "/profile.php":
        profile_id = parse_qs(parsed.query).get("id", [""])[0]
        if not profile_id:
            return []
        base = f"https://{host}/profile.php?id={profile_id}"
        variants.extend([
            f"{base}&sk=about_contact_and_basic_info",
            f"{base}&sk=about",
            base,
        ])
    else:
        base_path = parsed.path.rstrip("/")
        base = f"https://{host}{base_path}"
        variants.extend([
            f"{base}/about_contact_and_basic_info",
            f"{base}/about",
            base,
        ])

    out: list[str] = []
    seen: set[str] = set()
    for variant in variants:
        if variant not in seen:
            out.append(variant)
            seen.add(variant)
    return out


def sanitize_facebook_search_query(query: str) -> str:
    query = re.sub(r"\bsite:facebook\.com\b", " ", query or "", flags=re.IGNORECASE)
    query = query.replace('"', ' ').replace("'", ' ')
    return _normalize_space(query)


# ---------------------------------------------------------------------------
# Browser creation
# ---------------------------------------------------------------------------

def _is_windows() -> bool:
    return os.name == "nt"


def _unblock_windows_file(path: str | Path) -> None:
    if not _is_windows():
        return
    target = Path(path)
    if not target.exists():
        return
    powershell = shutil.which("powershell") or shutil.which("pwsh")
    if powershell:
        with suppress(Exception):
            subprocess.run(
                [powershell, "-NoProfile", "-ExecutionPolicy", "Bypass",
                 "-Command", "Unblock-File -LiteralPath $args[0]", str(target)],
                check=False, capture_output=True, text=True, timeout=10,
            )
    with suppress(Exception):
        subprocess.run(
            ["cmd", "/c", f'del /f /q "{target}:Zone.Identifier"'],
            check=False, capture_output=True, text=True, timeout=10,
        )


def _make_clean_driver_copy(src: str | Path) -> Path:
    source = Path(src)
    temp_dir = Path(tempfile.gettempdir()) / "fb_lead_driver_cache"
    temp_dir.mkdir(parents=True, exist_ok=True)
    stamp = int(time.time() * 1000)
    dest = temp_dir / f"{source.stem}_{stamp}{source.suffix}"
    dest.write_bytes(source.read_bytes())
    with suppress(Exception):
        os.chmod(dest, stat.S_IRUSR | stat.S_IWUSR | stat.S_IXUSR)
    _unblock_windows_file(dest)
    return dest


def _iter_chromedriver_candidates() -> list[Path]:
    candidates: list[Path] = []
    seen: set[str] = set()

    def add(p):
        if not p:
            return
        candidate = Path(p)
        if candidate.is_dir():
            return
        resolved = str(candidate)
        with suppress(Exception):
            resolved = str(candidate.resolve())
        if resolved in seen or not candidate.exists():
            return
        seen.add(resolved)
        candidates.append(candidate)

    exe_names = ["chromedriver.exe", "chromedriver"] if _is_windows() else ["chromedriver"]
    for env_name in ("SE_CHROMEDRIVER", "CHROMEDRIVER", "WEBDRIVER_CHROME_DRIVER"):
        add(os.environ.get(env_name))
    for base in (Path.cwd(), Path(__file__).resolve().parent):
        for exe_name in exe_names:
            add(base / exe_name)
    for path_entry in os.environ.get("PATH", "").split(os.pathsep):
        if not path_entry.strip():
            continue
        for exe_name in exe_names:
            add(Path(path_entry.strip()) / exe_name)
    cache_roots = [Path.home() / ".cache" / "selenium", Path.home() / ".wdm"]
    appdata = os.environ.get("APPDATA")
    localappdata = os.environ.get("LOCALAPPDATA")
    if appdata:
        cache_roots.append(Path(appdata) / "undetected_chromedriver")
    if localappdata:
        cache_roots.append(Path(localappdata) / "selenium")
    for root in cache_roots:
        if not root.exists():
            continue
        for pattern in ("chromedriver*.exe", "chromedriver*", "undetected_chromedriver*.exe"):
            with suppress(Exception):
                for match in root.rglob(pattern):
                    if match.is_file():
                        add(match)
    return candidates


def _download_chromedriver_with_webdriver_manager() -> Path | None:
    with suppress(Exception):
        from webdriver_manager.chrome import ChromeDriverManager
        downloaded = Path(ChromeDriverManager().install())
        if downloaded.is_file() and downloaded.name.lower().startswith("chromedriver"):
            return downloaded
        fallback_name = "chromedriver.exe" if _is_windows() else "chromedriver"
        fallback = downloaded.parent / fallback_name
        if fallback.exists():
            return fallback
    return None


def create_browser(
    headless: bool = False,
    user_data_dir: str = "",
    profile_directory: str = "",
    browser_binary: str = "",
):
    from selenium import webdriver
    from selenium.webdriver.chrome.service import Service as ChromeService

    chrome_args = [
        "--disable-blink-features=AutomationControlled",
        "--disable-infobars",
        "--disable-popup-blocking",
        "--disable-notifications",
        "--no-first-run",
        "--no-default-browser-check",
        "--window-size=1400,1000",
    ]
    if headless:
        chrome_args.append("--headless=new")

    def build_options():
        opts = webdriver.ChromeOptions()
        for arg in chrome_args:
            opts.add_argument(arg)
        opts.add_experimental_option("excludeSwitches", ["enable-automation"])
        opts.add_experimental_option("useAutomationExtension", False)
        if user_data_dir:
            opts.add_argument(f"--user-data-dir={user_data_dir}")
        if profile_directory:
            opts.add_argument(f"--profile-directory={profile_directory}")
        if browser_binary:
            opts.binary_location = browser_binary
        return opts

    def finalize(driver):
        with suppress(Exception):
            driver.execute_cdp_cmd(
                "Page.addScriptToEvaluateOnNewDocument",
                {"source": "Object.defineProperty(navigator, 'webdriver', {get: () => undefined});"},
            )
        driver.set_page_load_timeout(35)
        driver.set_script_timeout(35)
        driver.implicitly_wait(4)
        return driver

    errors: list[str] = []

    options = build_options()
    env_backup = {
        "SE_SKIP_DRIVER_IN_PATH": os.environ.get("SE_SKIP_DRIVER_IN_PATH"),
        "SE_CHROMEDRIVER": os.environ.get("SE_CHROMEDRIVER"),
    }
    try:
        # 1) Best path: built-in Selenium Manager, while ignoring broken PATH drivers.
        os.environ["SE_SKIP_DRIVER_IN_PATH"] = "true"
        os.environ.pop("SE_CHROMEDRIVER", None)
        try:
            return finalize(webdriver.Chrome(options=options))
        except Exception as exc:
            errors.append(f"Selenium Manager launch failed: {exc}")

        # 2) Local/cache drivers, but only after making a fresh temp copy.
        for candidate in _iter_chromedriver_candidates():
            try:
                clean_candidate = _make_clean_driver_copy(candidate)
                service = ChromeService(executable_path=str(clean_candidate))
                return finalize(webdriver.Chrome(service=service, options=options))
            except Exception as exc:
                errors.append(f"Candidate driver failed ({candidate}): {exc}")

        # 3) webdriver-manager download.
        downloaded = _download_chromedriver_with_webdriver_manager()
        if downloaded:
            try:
                clean_downloaded = _make_clean_driver_copy(downloaded)
                service = ChromeService(executable_path=str(clean_downloaded))
                return finalize(webdriver.Chrome(service=service, options=options))
            except Exception as exc:
                errors.append(f"webdriver-manager launch failed: {exc}")

    finally:
        for env_name, old_value in env_backup.items():
            if old_value is None:
                os.environ.pop(env_name, None)
            else:
                os.environ[env_name] = old_value

    message = "\n".join(f"- {m}" for m in errors[-10:])
    raise RuntimeError(f"Could not start Chrome.\n{message}")


# ---------------------------------------------------------------------------
# Facebook login / search helpers
# ---------------------------------------------------------------------------

def dismiss_facebook_dialogs(driver):
    with suppress(Exception):
        buttons = driver.find_elements("css selector", "button, div[role='button']")
        for button in buttons[:120]:
            label = _normalize_space(
                " ".join(
                    part for part in [
                        button.text or "",
                        button.get_attribute("aria-label") or "",
                    ] if part
                )
            ).lower()
            if any(
                token in label
                for token in (
                    "allow all cookies", "decline optional cookies",
                    "only allow essential cookies", "not now", "close", "ok",
                    "accept", "allow essential",
                )
            ):
                with suppress(Exception):
                    button.click()
                    time.sleep(0.4)


def facebook_login_page_detected(driver) -> bool:
    try:
        from selenium.webdriver.common.by import By
    except Exception:
        return False

    selectors = [
        (By.CSS_SELECTOR, "input[name='email']"),
        (By.CSS_SELECTOR, "input[name='pass']"),
        (By.CSS_SELECTOR, "form[action*='login']"),
        (By.XPATH, "//button[contains(., 'Log in') or contains(., 'Log In') or contains(., 'Log into Facebook')]"),
    ]
    for by, selector in selectors:
        with suppress(Exception):
            for element in driver.find_elements(by, selector):
                if element.is_displayed():
                    return True
    current_url = (getattr(driver, "current_url", "") or "").lower()
    return any(token in current_url for token in ("/login", "checkpoint", "recover"))


def facebook_auto_login(driver) -> bool:
    if not FB_EMAIL or not FB_PASSWORD:
        log("No FB_EMAIL / FB_PASSWORD in .env — manual login required")
        return False

    try:
        from selenium.webdriver.common.by import By
        from selenium.webdriver.common.keys import Keys
    except Exception:
        return False

    try:
        email_field = driver.find_element(By.CSS_SELECTOR, "input[name='email']")
        pass_field = driver.find_element(By.CSS_SELECTOR, "input[name='pass']")
        email_field.clear()
        email_field.send_keys(FB_EMAIL)
        time.sleep(0.3)
        pass_field.clear()
        pass_field.send_keys(FB_PASSWORD)
        time.sleep(0.3)
        pass_field.send_keys(Keys.ENTER)
        log("Auto-login submitted — waiting for Facebook to load...")
        time.sleep(5)
        dismiss_facebook_dialogs(driver)
        return True
    except Exception as exc:
        log(f"Auto-login failed: {exc}")
        return False


def find_facebook_search_bar(driver):
    try:
        from selenium.webdriver.common.by import By
    except Exception:
        return None

    input_selectors = [
        (By.CSS_SELECTOR, "input[placeholder*='Search Facebook']"),
        (By.CSS_SELECTOR, "input[aria-label*='Search Facebook']"),
        (By.CSS_SELECTOR, "input[type='search']"),
        (By.CSS_SELECTOR, "[role='search'] input"),
        (By.CSS_SELECTOR, "input[placeholder*='Search']"),
        (By.CSS_SELECTOR, "input[aria-label*='Search']"),
    ]
    for by, selector in input_selectors:
        with suppress(Exception):
            for element in driver.find_elements(by, selector):
                if element.is_displayed():
                    return element

    icon_selectors = [
        (By.CSS_SELECTOR, "[aria-label='Search Facebook']"),
        (By.CSS_SELECTOR, "[aria-label='Search']"),
        (By.XPATH, FACEBOOK_SEARCH_BAR_XPATH),
    ]
    for by, selector in icon_selectors:
        with suppress(Exception):
            for icon in driver.find_elements(by, selector):
                if not icon.is_displayed():
                    continue
                icon.click()
                time.sleep(0.8)
                for by2, selector2 in input_selectors:
                    with suppress(Exception):
                        for element in driver.find_elements(by2, selector2):
                            if element.is_displayed():
                                return element
                with suppress(Exception):
                    active = driver.switch_to.active_element
                    if (active.tag_name or "").lower() in {"input", "textarea"}:
                        return active
    return None


def ensure_facebook_search_ready(driver, timeout: int = FACEBOOK_MANUAL_LOGIN_WAIT_SECONDS) -> bool:
    if not driver:
        return False

    safe_driver_get(driver, "https://www.facebook.com/")
    dismiss_facebook_dialogs(driver)

    if facebook_login_page_detected(driver):
        facebook_auto_login(driver)

    deadline = time.time() + max(10, timeout)
    login_message_sent = False

    while time.time() < deadline:
        dismiss_facebook_dialogs(driver)
        element = find_facebook_search_bar(driver)
        if element is not None:
            return True

        if facebook_login_page_detected(driver) and not login_message_sent:
            log("Facebook login still required. Complete it in the browser window; the scraper will continue automatically.")
            login_message_sent = True

        time.sleep(1.0)

    return False


def _facebook_search_targets(driver, search_element) -> list:
    try:
        from selenium.webdriver.common.by import By
    except Exception:
        return []

    candidates: list = []
    if search_element is not None:
        candidates.append(search_element)
        with suppress(Exception):
            candidates.extend(search_element.find_elements(By.XPATH, ".//input | .//textarea"))

    with suppress(Exception):
        active = driver.switch_to.active_element
        if active is not None:
            tag = (active.tag_name or "").lower()
            if tag in ("input", "textarea") or active.get_attribute("contenteditable") == "true":
                candidates.append(active)

    selectors = [
        (By.CSS_SELECTOR, "input[placeholder*='Search Facebook']"),
        (By.CSS_SELECTOR, "input[aria-label*='Search Facebook']"),
        (By.CSS_SELECTOR, "input[type='search']"),
        (By.CSS_SELECTOR, "[role='search'] input"),
        (By.CSS_SELECTOR, "input[placeholder*='Search']"),
        (By.CSS_SELECTOR, "input[aria-label*='Search']"),
    ]
    for by, selector in selectors:
        with suppress(Exception):
            candidates.extend(driver.find_elements(by, selector))

    unique: list = []
    seen: set[str] = set()
    for candidate in candidates:
        if candidate is None:
            continue
        key = getattr(candidate, "id", None) or repr(candidate)
        if key in seen:
            continue
        seen.add(key)
        unique.append(candidate)
    return unique


def submit_facebook_home_search(driver, query: str) -> bool:
    query = sanitize_facebook_search_query(query)
    if not query:
        return False

    if not ensure_facebook_search_ready(driver):
        return False

    try:
        from selenium.webdriver.common.action_chains import ActionChains
        from selenium.webdriver.common.keys import Keys
    except Exception:
        return False

    search_element = find_facebook_search_bar(driver)
    if search_element is None:
        return False

    before_url = (getattr(driver, "current_url", "") or "").strip()
    with suppress(Exception):
        driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", search_element)
    time.sleep(0.25)

    try:
        ActionChains(driver).move_to_element(search_element).pause(0.1).click(search_element).perform()
    except Exception:
        with suppress(Exception):
            search_element.click()
    time.sleep(0.5)

    submitted = False
    for target in _facebook_search_targets(driver, search_element):
        try:
            target.send_keys(Keys.CONTROL, "a")
            target.send_keys(Keys.DELETE)
        except Exception:
            with suppress(Exception):
                target.clear()
        try:
            target.send_keys(query)
            time.sleep(0.25)
            target.send_keys(Keys.ENTER)
            submitted = True
            break
        except Exception:
            continue

    if not submitted:
        return False

    end = time.time() + FACEBOOK_SEARCH_WAIT_SECONDS
    while time.time() < end:
        current = (getattr(driver, "current_url", "") or "").strip()
        if current != before_url and "search" in current.lower():
            break
        time.sleep(0.5)

    dismiss_facebook_dialogs(driver)
    time.sleep(1.0)
    return True


# ---------------------------------------------------------------------------
# Search result extraction
# ---------------------------------------------------------------------------

def _candidate_score(query: str, name: str, description: str, url: str) -> int:
    query_tokens = {t for t in re.findall(r"[a-z0-9]+", (query or "").lower()) if len(t) > 2}
    text = " ".join([
        name or "",
        description or "",
        facebook_slug_to_label(url) or "",
    ]).lower()
    candidate_tokens = {t for t in re.findall(r"[a-z0-9]+", text) if len(t) > 2}

    score = 0
    overlap = len(query_tokens & candidate_tokens)
    score += overlap * 3

    if name:
        score += 3
    if description:
        score += 1
    if urlparse(url).path == "/profile.php":
        score -= 1
    if any(token in text for token in ("contractor", "construction", "builder", "company", "roof", "plumb", "electric")):
        score += 1
    return score


def _extract_facebook_candidates_from_articles(driver, query: str, max_results: int = 12) -> list[dict]:
    try:
        from selenium.webdriver.common.by import By
    except Exception:
        return []

    results: list[dict] = []
    seen: set[str] = set()

    with suppress(Exception):
        articles = driver.find_elements(By.CSS_SELECTOR, "div[role='article']")
    if not articles:
        articles = []

    for article in articles:
        page_url = ""
        name = ""
        description = ""

        with suppress(Exception):
            description = _normalize_space(article.text or "")

        with suppress(Exception):
            for element in article.find_elements(By.CSS_SELECTOR, "[aria-label]"):
                aria = (element.get_attribute("aria-label") or "").strip()
                if aria.startswith("Profile photo of "):
                    name = aria[len("Profile photo of "):].strip()
                    break

        with suppress(Exception):
            anchors = article.find_elements(By.CSS_SELECTOR, "a[href]")
        if not anchors:
            anchors = []

        for anchor in anchors:
            href = anchor.get_attribute("href") or ""
            normalized = normalize_facebook_page_url(decode_tracking_url(href))
            if not normalized:
                continue
            page_url = normalized
            if not name:
                candidate_name = normalize_facebook_title(
                    anchor.text or anchor.get_attribute("aria-label") or ""
                )
                if candidate_name and candidate_name.lower() not in {
                    "facebook", "follow", "following", "share", "see more",
                    "more", "pages", "results",
                }:
                    name = candidate_name
            break

        if not page_url or page_url in seen:
            continue

        if not name:
            name = normalize_facebook_title(facebook_slug_to_label(page_url))

        score = _candidate_score(query, name, description, page_url)
        results.append({
            "url": page_url,
            "name": name,
            "description": description,
            "score": score,
        })
        seen.add(page_url)

    results.sort(key=lambda item: (item["score"], len(item["name"] or ""), len(item["description"] or "")), reverse=True)
    return results[:max_results]


def _extract_facebook_candidates_from_anchors(driver, query: str, max_results: int = 12) -> list[dict]:
    try:
        from selenium.webdriver.common.by import By
    except Exception:
        return []

    results: list[dict] = []
    seen: set[str] = set()

    with suppress(Exception):
        anchors = driver.find_elements(By.CSS_SELECTOR, "a[href]")
    if not anchors:
        anchors = []

    for anchor in anchors:
        href = anchor.get_attribute("href") or ""
        page_url = normalize_facebook_page_url(decode_tracking_url(href))
        if not page_url or page_url in seen:
            continue
        name = normalize_facebook_title(anchor.text or anchor.get_attribute("aria-label") or "")
        if not name:
            name = normalize_facebook_title(facebook_slug_to_label(page_url))
        description = ""
        score = _candidate_score(query, name, description, page_url)
        if score < 2:
            continue
        results.append({
            "url": page_url,
            "name": name,
            "description": description,
            "score": score,
        })
        seen.add(page_url)

    results.sort(key=lambda item: (item["score"], len(item["name"] or "")), reverse=True)
    return results[:max_results]


def _wait_for_facebook_results(driver, timeout: int = 12) -> bool:
    try:
        from selenium.webdriver.common.by import By
        from selenium.webdriver.support.ui import WebDriverWait
    except Exception:
        return False

    selectors = [
        (By.CSS_SELECTOR, "div[role='article']"),
        (By.CSS_SELECTOR, "a[href*='/search/pages']"),
        (By.CSS_SELECTOR, "a[href*='facebook.com/']"),
    ]
    end = time.time() + timeout
    while time.time() < end:
        dismiss_facebook_dialogs(driver)
        for by, selector in selectors:
            with suppress(Exception):
                elements = driver.find_elements(by, selector)
                if any(getattr(e, "is_displayed", lambda: True)() for e in elements):
                    return True
        time.sleep(0.5)
    return False


def search_facebook_page_candidates_via_ui(
    driver,
    query: str,
    max_results: int = 8,
    debug_dir: str = "",
) -> list[dict]:
    if not driver:
        return []

    cleaned_query = sanitize_facebook_search_query(query)
    if not cleaned_query:
        return []

    # Make sure we can reach Facebook home and either auto-login or let the user login.
    current = (getattr(driver, "current_url", "") or "").lower()
    if "facebook.com" not in current:
        safe_driver_get(driver, "https://www.facebook.com/")
        time.sleep(1.2)

    if facebook_login_page_detected(driver):
        log("  Logging into Facebook...")
        facebook_auto_login(driver)
        time.sleep(2.0)

    if facebook_login_page_detected(driver):
        log("  Manual Facebook login required — complete it in the browser window.")
        deadline = time.time() + FACEBOOK_MANUAL_LOGIN_WAIT_SECONDS
        while time.time() < deadline and facebook_login_page_detected(driver):
            time.sleep(1.0)
        if facebook_login_page_detected(driver):
            log("  Login timeout — continuing with external-search fallback only.")
            return []

    # Primary route: direct Pages search URL.
    pages_url = f"https://www.facebook.com/search/pages/?q={quote_plus(cleaned_query)}"
    log(f"  Navigating directly to: {pages_url}")
    if not safe_driver_get(driver, pages_url):
        log("  Direct navigation failed; trying home-page search fallback...")
        if not submit_facebook_home_search(driver, cleaned_query):
            return []

    dismiss_facebook_dialogs(driver)
    _wait_for_facebook_results(driver, timeout=12)

    candidates: list[dict] = []
    seen: set[str] = set()

    for scroll_round in range(max(2, FACEBOOK_SEARCH_SCROLL_ROUNDS)):
        current_page_candidates = _extract_facebook_candidates_from_articles(driver, cleaned_query, max_results=max_results * 3)
        if not current_page_candidates:
            current_page_candidates = _extract_facebook_candidates_from_anchors(driver, cleaned_query, max_results=max_results * 3)

        for candidate in current_page_candidates:
            url = candidate.get("url", "")
            if not url or url in seen:
                continue
            seen.add(url)
            candidates.append(candidate)
            if len(candidates) >= max_results:
                return candidates[:max_results]

        with suppress(Exception):
            driver.execute_script(
                "window.scrollBy(0, Math.max(700, Math.floor(window.innerHeight * 0.85)));"
            )
        time.sleep(1.0)
        dismiss_facebook_dialogs(driver)

    if debug_dir:
        _save_debug_html(debug_dir, f"fb_search_{cleaned_query}", driver.page_source or "")
        _save_debug_png(driver, debug_dir, f"fb_search_{cleaned_query}")

    return candidates[:max_results]


# ---------------------------------------------------------------------------
# External search fallback
# ---------------------------------------------------------------------------

def search_duckduckgo(session: requests.Session, query: str, max_results: int = 8) -> list[str]:
    url = f"https://html.duckduckgo.com/html/?q={quote_plus(query)}"
    html_text = fetch_html(session, url, timeout=20)
    if not html_text:
        return []
    soup = make_soup(html_text)
    urls: list[str] = []
    seen: set[str] = set()
    for anchor in soup.select("a.result__a, a.result-link, h2.result__title a, a[data-testid='result-title-a']"):
        href = decode_tracking_url(anchor.get("href", ""))
        cleaned = clean_url(href)
        if cleaned and cleaned not in seen:
            urls.append(cleaned)
            seen.add(cleaned)
        if len(urls) >= max_results:
            break
    return urls


def search_bing(session: requests.Session, query: str, max_results: int = 8) -> list[str]:
    url = f"https://www.bing.com/search?q={quote_plus(query)}"
    html_text = fetch_html(session, url, timeout=20)
    if not html_text:
        return []
    soup = make_soup(html_text)
    urls: list[str] = []
    seen: set[str] = set()
    for anchor in soup.select("li.b_algo h2 a, h2 a"):
        href = decode_tracking_url(anchor.get("href", ""))
        cleaned = clean_url(href)
        if cleaned and cleaned not in seen:
            urls.append(cleaned)
            seen.add(cleaned)
        if len(urls) >= max_results:
            break
    return urls


def search_facebook_page_candidates(
    session: requests.Session,
    query: str,
    max_results: int = 8,
    driver=None,
    debug_dir: str = "",
) -> list[dict]:
    cleaned_query = sanitize_facebook_search_query(query)
    candidates: list[dict] = []
    seen: set[str] = set()

    if driver is not None:
        with suppress(Exception):
            ui_results = search_facebook_page_candidates_via_ui(driver, cleaned_query, max_results=max_results, debug_dir=debug_dir)
            for result in ui_results:
                url = normalize_facebook_page_url(result.get("url", ""))
                if not url or url in seen:
                    continue
                result["url"] = url
                candidates.append(result)
                seen.add(url)
                if len(candidates) >= max_results:
                    return candidates

    # External fallback. It is slower and less precise, so it stays secondary.
    search_query = cleaned_query
    if "facebook.com" not in search_query.lower():
        search_query = f'site:facebook.com "{cleaned_query}"'

    for search_fn in (search_duckduckgo, search_bing):
        with suppress(Exception):
            urls = search_fn(session, search_query, max_results=max_results * 2)
            for url in urls:
                normalized = normalize_facebook_page_url(url)
                if not normalized or normalized in seen:
                    continue
                candidates.append({
                    "url": normalized,
                    "name": normalize_facebook_title(facebook_slug_to_label(normalized)),
                    "description": "",
                    "score": 0,
                })
                seen.add(normalized)
                if len(candidates) >= max_results:
                    return candidates

    return candidates


# ---------------------------------------------------------------------------
# Contact extraction
# ---------------------------------------------------------------------------

def _lxml_node_text(node) -> str:
    if node is None:
        return ""
    if isinstance(node, str):
        return _normalize_space(node)
    with suppress(Exception):
        return _normalize_space(" ".join(node.itertext()))
    with suppress(Exception):
        return _normalize_space(etree.tostring(node, method="text", encoding="unicode"))
    return _normalize_space(str(node))


def extract_primary_external_website(soup: BeautifulSoup) -> str:
    scored: list[tuple[int, str]] = []
    seen: set[str] = set()

    for anchor in soup.select("a[href]"):
        raw_href = anchor.get("href", "")
        href = decode_tracking_url(raw_href)
        cleaned = clean_url(href)
        if not cleaned or cleaned in seen:
            continue
        seen.add(cleaned)
        if _is_blocked_website(cleaned):
            continue
        anchor_text = (anchor.get_text(" ", strip=True) or "").lower()
        score = 0
        if "/l.php" in raw_href or "l.facebook.com" in raw_href:
            score += 3
        if any(t in anchor_text for t in ("website", "visit", "official", "site", "home", "www", "http")):
            score += 2
        if cleaned.startswith("https://"):
            score += 1
        scored.append((score, cleaned))

    if not scored:
        return ""
    scored.sort(key=lambda item: (item[0], len(item[1])), reverse=True)
    return scored[0][1]


def extract_contact_snapshot(soup: BeautifulSoup, source_url: str = "") -> dict:
    title_candidates: list[str] = []

    for meta in soup.select("meta[property='og:title'], meta[name='twitter:title']"):
        cleaned = normalize_facebook_title(meta.get("content", ""))
        if cleaned:
            title_candidates.append(cleaned)

    h1 = soup.find(["h1", "h2"])
    if h1:
        cleaned = normalize_facebook_title(h1.get_text(" ", strip=True))
        if cleaned:
            title_candidates.append(cleaned)

    if source_url:
        slug = normalize_facebook_title(facebook_slug_to_label(source_url))
        if slug:
            title_candidates.append(slug)

    if soup.title:
        cleaned = normalize_facebook_title(soup.title.get_text(" ", strip=True))
        if cleaned:
            title_candidates.append(cleaned)

    texts: list[str] = []
    for meta in soup.select("meta[property='og:description'], meta[name='description'], meta[name='twitter:description']"):
        content = meta.get("content", "").strip()
        if content:
            texts.append(content)

    page_text = soup.get_text(" ", strip=True)
    if page_text:
        texts.append(page_text)

    blob = " | ".join(part for part in texts if part)
    phones = extract_phones(blob)
    emails = extract_emails(blob)
    company_name = first_non_empty(*title_candidates)
    website = extract_primary_external_website(soup)

    return {
        "company_name": company_name,
        "phones": phones,
        "emails": emails,
        "website": website,
        "blob": blob[:12000],
    }


def extract_graphql_contacts_from_text(html_text: str) -> dict:
    result = {"phones": [], "emails": [], "website": "", "hit": False}
    if not html_text:
        return result

    seen_phones: set[str] = set()
    seen_emails: set[str] = set()

    for match in re.finditer(r'"pressable_profile_field_type"\s*:\s*"(BUSINESS_PHONE|BUSINESS_EMAIL|BUSINESS_WEBSITE)"', html_text):
        field_type = match.group(1)
        context_start = max(0, match.start() - 3000)
        context_end = min(len(html_text), match.end() + 3000)
        context = html_text[context_start:context_end]
        text_match = re.search(r'"item_subtitle"\s*:\s*\{"text"\s*:\s*\{"text"\s*:\s*"([^"]+)"', context)
        if not text_match:
            continue
        text_val = text_match.group(1)
        with suppress(Exception):
            text_val = text_val.encode().decode("unicode_escape")
        result["hit"] = True

        if field_type == "BUSINESS_PHONE":
            for phone in extract_phones(text_val):
                normalized = normalize_phone(phone)
                if normalized and normalized not in seen_phones:
                    seen_phones.add(normalized)
                    result["phones"].append(normalized)
        elif field_type == "BUSINESS_EMAIL":
            for email in extract_emails(text_val):
                if email not in seen_emails:
                    seen_emails.add(email)
                    result["emails"].append(email)
        elif field_type == "BUSINESS_WEBSITE" and not result["website"]:
            cleaned = clean_url(decode_tracking_url(text_val))
            if cleaned and not _is_blocked_website(cleaned):
                result["website"] = cleaned

    return result


def _fb_heuristic_extract_from_driver(driver) -> dict:
    try:
        from selenium.webdriver.common.by import By
    except Exception:
        return {"phones": [], "emails": [], "website": "", "hit": False}

    phones: list[str] = []
    emails: list[str] = []
    website = ""
    hit = False
    seen_phones: set[str] = set()
    seen_emails: set[str] = set()

    def add_phone(value: str):
        nonlocal hit
        for phone in extract_phones(value):
            normalized = normalize_phone(phone)
            if normalized and normalized not in seen_phones:
                seen_phones.add(normalized)
                phones.append(normalized)
                hit = True

    def add_email(value: str):
        nonlocal hit
        for email in extract_emails(value):
            if email not in seen_emails:
                seen_emails.add(email)
                emails.append(email)
                hit = True

    with suppress(Exception):
        body = driver.find_element(By.TAG_NAME, "body")
        body_text = body.text or body.get_attribute("innerText") or body.get_attribute("textContent") or ""
        add_phone(body_text)
        add_email(body_text)

    with suppress(Exception):
        for a_tag in driver.find_elements(By.CSS_SELECTOR, "a[href^='mailto:']"):
            href = (a_tag.get_attribute("href") or "").replace("mailto:", "").split("?")[0].strip()
            add_email(href)

    with suppress(Exception):
        for a_tag in driver.find_elements(By.CSS_SELECTOR, "a[href^='tel:']"):
            href = (a_tag.get_attribute("href") or "").replace("tel:", "").strip()
            add_phone(href)

    with suppress(Exception):
        for elem in driver.find_elements(By.CSS_SELECTOR, "[aria-label]"):
            label = (elem.get_attribute("aria-label") or "").lower()
            text = _normalize_space(" ".join([
                elem.text or "",
                elem.get_attribute("innerText") or "",
                elem.get_attribute("textContent") or "",
                label,
            ]))
            if any(token in label for token in _FB_PHONE_LABELS):
                add_phone(text)
            if any(token in label for token in _FB_EMAIL_LABELS):
                add_email(text)
            if any(token in label for token in _FB_WEBSITE_LABELS):
                with suppress(Exception):
                    for a_tag in elem.find_elements(By.CSS_SELECTOR, "a[href]"):
                        href = clean_url(decode_tracking_url(a_tag.get_attribute("href") or ""))
                        if href and not _is_blocked_website(href):
                            website = href
                            hit = True
                            break

    with suppress(Exception):
        for span in driver.find_elements(By.TAG_NAME, "span"):
            span_text = (span.text or "").strip().lower()
            if not span_text:
                continue

            field_type = None
            if any(label in span_text for label in _FB_PHONE_LABELS):
                field_type = "phone"
            elif any(label in span_text for label in _FB_EMAIL_LABELS):
                field_type = "email"
            elif any(label in span_text for label in _FB_WEBSITE_LABELS):
                field_type = "website"

            if field_type is None:
                continue

            for level in _FB_ANCESTOR_LEVELS:
                try:
                    container = span.find_element(By.XPATH, "/".join([".."] * level))
                    text = container.text or container.get_attribute("innerText") or container.get_attribute("textContent") or ""
                    if field_type == "phone":
                        add_phone(text)
                    elif field_type == "email":
                        add_email(text)
                    else:
                        for a_tag in container.find_elements(By.CSS_SELECTOR, "a[href]"):
                            href = clean_url(decode_tracking_url(a_tag.get_attribute("href") or ""))
                            if href and not _is_blocked_website(href):
                                website = href
                                hit = True
                                break
                except Exception:
                    continue

    return {"phones": phones, "emails": emails, "website": website, "hit": hit}


def extract_facebook_xpath_contacts_from_html(html_text: str) -> dict:
    data = {"phones": [], "emails": [], "raw_phone_text": "", "raw_email_text": "", "email_href": "", "xpath_hit": False}
    if not html_text or etree is None:
        return data

    try:
        tree = etree.HTML(html_text)
    except Exception:
        return data

    raw_phone_texts: list[str] = []
    for node in tree.xpath(FACEBOOK_PHONE_XPATH):
        node_text = _lxml_node_text(node)
        if node_text:
            raw_phone_texts.append(node_text)
    raw_phone_texts = _dedupe_ci(raw_phone_texts)
    if raw_phone_texts:
        data["raw_phone_text"] = " | ".join(raw_phone_texts)
        data["phones"] = _merge_phone_lists(extract_phones(" | ".join(raw_phone_texts)))
        data["xpath_hit"] = True

    raw_email_texts: list[str] = []
    href_emails: list[str] = []
    href_values: list[str] = []
    for node in tree.xpath(FACEBOOK_EMAIL_XPATH):
        node_text = _lxml_node_text(node)
        if node_text:
            raw_email_texts.append(node_text)
        href = ""
        with suppress(Exception):
            href = (node.get("href") or "").strip()
        if href:
            href_values.append(href)
            if href.lower().startswith("mailto:"):
                href_emails.extend(extract_emails(href.replace("mailto:", "", 1)))

    raw_email_texts = _dedupe_ci(raw_email_texts)
    href_values = _dedupe_ci(href_values)
    if raw_email_texts:
        data["raw_email_text"] = " | ".join(raw_email_texts)
        data["xpath_hit"] = True
    if href_values:
        data["email_href"] = " | ".join(href_values)
        data["xpath_hit"] = True

    text_emails = extract_emails(" | ".join(raw_email_texts))
    data["emails"] = _merge_email_lists(href_emails, text_emails)
    return data


def extract_facebook_xpath_contacts_from_driver_dom(driver) -> dict:
    data = {"phones": [], "emails": [], "raw_phone_text": "", "raw_email_text": "", "email_href": "", "xpath_hit": False}
    try:
        from selenium.webdriver.common.by import By
    except Exception:
        return data

    raw_phone_texts: list[str] = []
    raw_email_texts: list[str] = []
    href_emails: list[str] = []
    href_values: list[str] = []

    with suppress(Exception):
        for element in driver.find_elements(By.XPATH, FACEBOOK_PHONE_XPATH):
            text_value = _normalize_space(element.text or element.get_attribute("innerText") or element.get_attribute("textContent") or "")
            if text_value:
                raw_phone_texts.append(text_value)

    with suppress(Exception):
        for element in driver.find_elements(By.XPATH, FACEBOOK_EMAIL_XPATH):
            text_value = _normalize_space(element.text or element.get_attribute("innerText") or element.get_attribute("textContent") or "")
            href = (element.get_attribute("href") or "").strip()
            if text_value:
                raw_email_texts.append(text_value)
            if href:
                href_values.append(href)
                if href.lower().startswith("mailto:"):
                    href_emails.extend(extract_emails(href.replace("mailto:", "", 1)))

    raw_phone_texts = _dedupe_ci(raw_phone_texts)
    raw_email_texts = _dedupe_ci(raw_email_texts)
    href_values = _dedupe_ci(href_values)

    if raw_phone_texts:
        data["raw_phone_text"] = " | ".join(raw_phone_texts)
        data["phones"] = _merge_phone_lists(extract_phones(" | ".join(raw_phone_texts)))
        data["xpath_hit"] = True
    if raw_email_texts:
        data["raw_email_text"] = " | ".join(raw_email_texts)
        data["xpath_hit"] = True
    if href_values:
        data["email_href"] = " | ".join(href_values)
        data["xpath_hit"] = True

    text_emails = extract_emails(" | ".join(raw_email_texts))
    data["emails"] = _merge_email_lists(href_emails, text_emails)
    return data


def fetch_facebook_variant_snapshot(driver, variant_url: str, timeout: int = 16, debug_dir: str = "") -> dict:
    snapshot = {
        "company_name": "",
        "phones": [],
        "emails": [],
        "website": "",
        "blob": "",
        "source_url": variant_url,
        "methods": [],
    }
    if not driver or not variant_url:
        return snapshot

    if not safe_driver_get(driver, variant_url):
        return snapshot

    try:
        from selenium.webdriver.common.by import By
        from selenium.webdriver.support.ui import WebDriverWait
        from selenium.webdriver.support import expected_conditions as EC
        with suppress(Exception):
            WebDriverWait(driver, timeout).until(EC.presence_of_element_located((By.TAG_NAME, "body")))
    except Exception:
        pass

    time.sleep(FACEBOOK_DETAIL_WAIT_SECONDS)
    dismiss_facebook_dialogs(driver)

    for scroll_y in (0, 300, 800, 1400):
        with suppress(Exception):
            driver.execute_script(f"window.scrollTo(0, {scroll_y});")
        time.sleep(0.35)

    html_text = ""
    with suppress(Exception):
        html_text = driver.page_source or ""
    if not html_text:
        return snapshot

    soup = make_soup(html_text)
    generic = extract_contact_snapshot(soup, source_url=variant_url)
    graphql = extract_graphql_contacts_from_text(html_text)
    heuristic = _fb_heuristic_extract_from_driver(driver)
    xpath_html = extract_facebook_xpath_contacts_from_html(html_text)
    xpath_dom = extract_facebook_xpath_contacts_from_driver_dom(driver)

    snapshot["company_name"] = generic.get("company_name", "")
    snapshot["website"] = (
        graphql.get("website", "")
        or heuristic.get("website", "")
        or generic.get("website", "")
    )
    snapshot["phones"] = _merge_phone_lists(
        graphql.get("phones", []),
        heuristic.get("phones", []),
        xpath_html.get("phones", []),
        xpath_dom.get("phones", []),
        generic.get("phones", []),
    )
    snapshot["emails"] = _merge_email_lists(
        graphql.get("emails", []),
        heuristic.get("emails", []),
        xpath_html.get("emails", []),
        xpath_dom.get("emails", []),
        generic.get("emails", []),
    )
    snapshot["blob"] = generic.get("blob", "")

    if graphql.get("hit"):
        snapshot["methods"].append("GraphQL")
    if heuristic.get("hit"):
        snapshot["methods"].append("Heuristic")
    if xpath_html.get("xpath_hit") or xpath_dom.get("xpath_hit"):
        snapshot["methods"].append("XPath")
    if generic.get("phones") or generic.get("emails") or generic.get("website"):
        snapshot["methods"].append("Generic")

    if debug_dir:
        hint = facebook_slug_to_label(variant_url) or "facebook_variant"
        _save_debug_html(debug_dir, f"{hint}_{urlparse(variant_url).path.strip('/').replace('/', '_')}", html_text)

    return snapshot


def fetch_public_facebook_page(
    session: requests.Session,
    page_url: str,
    timeout: int = 16,
    driver=None,
    debug_dir: str = "",
) -> dict:
    normalized = normalize_facebook_page_url(page_url)
    if not normalized:
        return {}

    variants = build_facebook_page_variants(normalized)
    if normalized not in variants:
        variants.append(normalized)

    merged_company = ""
    merged_phones: list[str] = []
    merged_emails: list[str] = []
    merged_website = ""
    notes_parts: list[str] = []
    used_urls: list[str] = []

    # IMPORTANT FIX:
    # The original script built About-page URLs but never used them. That meant
    # your absolute email/phone XPaths were never visited. We now visit:
    #   about_contact_and_basic_info -> about -> main page
    for variant in variants:
        snapshot = fetch_facebook_variant_snapshot(driver, variant, timeout=timeout, debug_dir=debug_dir) if driver else {}
        if not snapshot:
            continue

        if not merged_company and snapshot.get("company_name"):
            merged_company = snapshot["company_name"]
        merged_phones = _merge_phone_lists(merged_phones, snapshot.get("phones", []))
        merged_emails = _merge_email_lists(merged_emails, snapshot.get("emails", []))
        if not merged_website and snapshot.get("website"):
            merged_website = snapshot["website"]
        used_urls.append(variant)

        if snapshot.get("methods"):
            notes_parts.append(f"{variant} -> " + ",".join(snapshot["methods"]))

        # If we already have enough contact info, stop early.
        if merged_phones and merged_emails and merged_website:
            break

    if not merged_company:
        merged_company = facebook_slug_to_label(normalized)

    # Optional HTTP fallback for name/website only.
    if session and (not merged_company or not merged_website):
        for variant in variants:
            html_text = fetch_html(session, variant, timeout=12)
            if not html_text:
                continue
            generic = extract_contact_snapshot(make_soup(html_text), source_url=variant)
            if not merged_company and generic.get("company_name"):
                merged_company = generic["company_name"]
            if not merged_website and generic.get("website"):
                merged_website = generic["website"]
            if merged_company and merged_website:
                break

    # Skip empty leads unless we captured something useful.
    if not any([merged_phones, merged_emails, merged_website]):
        return {}

    if len(merged_phones) > 1:
        notes_parts.append("Alt phones: " + ", ".join(merged_phones[1:]))
    if len(merged_emails) > 1:
        notes_parts.append("Alt emails: " + ", ".join(merged_emails[1:]))

    return {
        "facebook_page": normalized,
        "company_name": merged_company,
        "phone": merged_phones[0] if merged_phones else "",
        "email": merged_emails[0] if merged_emails else "",
        "website": merged_website,
        "blob": "",
        "notes": " | ".join(_dedupe_ci(notes_parts)),
        "checked_urls": used_urls,
    }


# ---------------------------------------------------------------------------
# Main scraper
# ---------------------------------------------------------------------------

def scrape_facebook_pages(
    search_terms: str | Iterable[str],
    locations: list[tuple[str, str]],
    headless: bool = False,
    results_per_location: int = DEFAULT_RESULTS_PER_CITY,
    candidate_search_depth: int = DEFAULT_CANDIDATE_SEARCH_DEPTH,
    user_data_dir: str = "",
    profile_directory: str = "",
    browser_binary: str = "",
    debug_dir: str = "",
    stop_event: threading.Event | None = None,
) -> list[Lead]:
    all_leads: list[Lead] = []
    driver = None
    session = build_session()
    profile_cache: dict[str, dict] = {}
    normalized_locations = dedupe_location_pairs(locations)
    keywords = normalize_search_terms(search_terms)

    results_per_location = max(1, int(results_per_location or DEFAULT_RESULTS_PER_CITY))
    candidate_search_depth = max(results_per_location * 3, int(candidate_search_depth or DEFAULT_CANDIDATE_SEARCH_DEPTH))

    log(
        "Facebook Pages -> starting "
        f"(keywords: {len(keywords)}, locations: {len(normalized_locations)}, "
        f"target leads/city: {results_per_location}, candidate depth/city: {candidate_search_depth})"
    )

    try:
        try:
            log("Facebook Pages -> launching browser...")
            driver = create_browser(
                headless=headless,
                user_data_dir=user_data_dir,
                profile_directory=profile_directory,
                browser_binary=browser_binary,
            )
        except Exception as exc:
            log(f"Facebook Pages -> browser unavailable: {exc}")
            log("Facebook Pages -> falling back to external search only (lower quality)")
            driver = None

        for city, state in normalized_locations:
            if stop_event and stop_event.is_set():
                log("Stop requested — ending after current city.")
                break

            log(f"\n{'=' * 60}")
            log(f"Facebook Pages -> city: {city}, {state}")
            log(f"{'=' * 60}")

            city_leads: list[Lead] = []
            seen_pages_this_city: set[str] = set()
            seen_company_keys: set[str] = set()

            for keyword_index, keyword in enumerate(keywords, start=1):
                if stop_event and stop_event.is_set():
                    break
                if len(city_leads) >= results_per_location:
                    break

                log(f"Keyword {keyword_index}/{len(keywords)} for {city}, {state}: '{keyword}'")
                queries = build_location_queries(keyword, city, state)
                candidates: list[dict] = []
                seen_candidate_urls: set[str] = set()

                per_query_limit = max(
                    results_per_location * 2,
                    min(candidate_search_depth, (candidate_search_depth // max(1, len(queries))) + results_per_location),
                )

                for query in queries:
                    if stop_event and stop_event.is_set():
                        break
                    remaining_capacity = candidate_search_depth - len(candidates)
                    if remaining_capacity <= 0:
                        break

                    log(f"  Query: '{query}'")
                    query_results = search_facebook_page_candidates(
                        session,
                        query,
                        max_results=min(remaining_capacity, per_query_limit),
                        driver=driver,
                        debug_dir=debug_dir,
                    )
                    log(f"    -> {len(query_results)} candidate page(s) returned")

                    for result in query_results:
                        url = normalize_facebook_page_url(result.get("url", ""))
                        if not url or url in seen_candidate_urls:
                            continue
                        result["url"] = url
                        candidates.append(result)
                        seen_candidate_urls.add(url)
                        if len(candidates) >= candidate_search_depth:
                            break

                    if len(candidates) >= candidate_search_depth:
                        break

                if not candidates:
                    log("  No candidates found for this keyword.")
                    continue

                log(
                    f"  Collected {len(candidates)} unique candidate page(s) for keyword '{keyword}' in {city}, {state}; "
                    f"targeting {results_per_location - len(city_leads)} more lead(s)"
                )

                for i, result in enumerate(candidates, 1):
                    if stop_event and stop_event.is_set():
                        break
                    if len(city_leads) >= results_per_location:
                        break

                    page_url = result.get("url", "")
                    search_name = _normalize_space(result.get("name", ""))
                    normalized_page = normalize_facebook_page_url(page_url)
                    if not normalized_page or normalized_page in seen_pages_this_city:
                        continue

                    log(f"  [{i}/{len(candidates)}] Fetching: {normalized_page}")
                    if search_name:
                        log(f"    Search result name: {search_name}")

                    if normalized_page in profile_cache:
                        profile = profile_cache[normalized_page]
                        log("    -> using cached profile")
                    else:
                        try:
                            profile = fetch_public_facebook_page(
                                session,
                                normalized_page,
                                timeout=18,
                                driver=driver,
                                debug_dir=debug_dir,
                            )
                        except Exception as exc:
                            log(f"    -> error fetching page: {exc}")
                            profile = None
                        profile_cache[normalized_page] = profile or {}

                    if not profile:
                        log("    -> no usable contact data extracted")
                        continue

                    company_name = (
                        search_name
                        or _normalize_space(profile.get("company_name", ""))
                        or facebook_slug_to_label(normalized_page)
                    )
                    if not company_name:
                        log("    -> skipped: no company name")
                        continue

                    company_key = normalize_name(company_name)
                    if company_key and company_key in seen_company_keys:
                        log("    -> skipped: duplicate company for this city")
                        continue

                    lead = Lead(
                        company_name=company_name,
                        phone=profile.get("phone", ""),
                        email=profile.get("email", ""),
                        city=city,
                        state=state,
                        website=profile.get("website", ""),
                        facebook_page=profile.get("facebook_page", normalized_page),
                        source="Facebook Pages",
                        search_keyword=keyword,
                        notes=profile.get("notes", ""),
                    )

                    if not any([lead.phone, lead.email, lead.website]):
                        log("    -> skipped: page had no phone, email, or website")
                        continue

                    all_leads.append(lead)
                    city_leads.append(lead)
                    seen_pages_this_city.add(normalized_page)
                    if company_key:
                        seen_company_keys.add(company_key)

                    log(
                        f"    -> LEAD {len(city_leads)}/{results_per_location} for {city}, {state}: "
                        f"{lead.company_name}"
                    )
                    if lead.phone:
                        log(f"       Phone: {lead.phone}")
                    if lead.email:
                        log(f"       Email: {lead.email}")
                    if lead.website:
                        log(f"       Website: {lead.website}")
                    log(f"       FB Page: {lead.facebook_page}")

                if len(city_leads) >= results_per_location:
                    log(f"Reached target for {city}, {state} after keyword '{keyword}'.")

            log(f"Completed {city}, {state}: {len(city_leads)}/{results_per_location} lead(s) collected")
            if len(city_leads) < results_per_location:
                log(
                    f"  Short by {results_per_location - len(city_leads)} lead(s). "
                    f"Add more keywords or increase candidate depth if needed."
                )

            time.sleep(random.uniform(0.8, 1.6))

    finally:
        if driver:
            with suppress(Exception):
                driver.quit()

    leads = deduplicate_leads(all_leads)
    log(f"\n{'=' * 60}")
    log(f"DONE — collected {len(leads)} unique leads total")
    log(f"{'=' * 60}\n")
    return leads


def print_results(leads: list[Lead]):
    if not leads:
        log("No leads found.")
        return

    grouped: dict[tuple[str, str], list[Lead]] = defaultdict(list)
    for lead in leads:
        grouped[(lead.city, lead.state)].append(lead)

    log(f"\n{'=' * 70}")
    log(f"  RESULTS: {len(leads)} leads")
    log(f"{'=' * 70}")

    for city_state, city_leads in grouped.items():
        city, state = city_state
        print(f"\n### {city}, {state} — {len(city_leads)} lead(s) ###")
        for i, lead in enumerate(city_leads, 1):
            print(f"\n--- Lead #{i} ---")
            print(f"  Company:  {lead.company_name}")
            print(f"  Phone:    {lead.phone or '—'}")
            print(f"  Email:    {lead.email or '—'}")
            print(f"  Website:  {lead.website or '—'}")
            print(f"  FB Page:  {lead.facebook_page or '—'}")
            print(f"  Location: {lead.city}, {lead.state}")
            print(f"  Keyword:  {lead.search_keyword or '—'}")
            if lead.notes:
                print(f"  Notes:    {lead.notes}")

    with_phone = sum(1 for l in leads if l.phone)
    with_email = sum(1 for l in leads if l.email)
    with_website = sum(1 for l in leads if l.website)
    print("\nSummary by city:")
    for city, state in dedupe_location_pairs(grouped.keys()):
        city_leads = grouped[(city, state)]
        city_phone = sum(1 for l in city_leads if l.phone)
        city_email = sum(1 for l in city_leads if l.email)
        city_website = sum(1 for l in city_leads if l.website)
        print(
            f"  {city}, {state}: {len(city_leads)} lead(s) | "
            f"{city_phone} with phone | {city_email} with email | {city_website} with website"
        )

    print(
        f"\nOverall: {len(leads)} leads | {with_phone} with phone | "
        f"{with_email} with email | {with_website} with website"
    )


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

EXCEL_HEADER_FILL = PatternFill(fill_type="solid", start_color="1F4E78", end_color="1F4E78")
EXCEL_HEADER_FONT = Font(color="FFFFFF", bold=True)
EXCEL_BODY_FONT = Font(name="Calibri", size=11)


def build_timestamped_excel_path(output_dir: str = "", prefix: str = "facebook_leads") -> Path:
    directory = Path(output_dir).expanduser() if output_dir else Path.cwd()
    directory.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return directory / f"{prefix}_{timestamp}.xlsx"


def export_leads_to_excel(
    leads: list[Lead],
    keywords: Iterable[str],
    locations: Iterable[tuple[str, str]],
    filepath: str | Path,
):
    filepath = Path(filepath)
    filepath.parent.mkdir(parents=True, exist_ok=True)
    leads = deduplicate_leads(leads)

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"

    columns = [
        ("Company Name", 32),
        ("Phone", 18),
        ("Email", 30),
        ("Website", 40),
        ("Facebook Page", 45),
        ("City", 18),
        ("State", 10),
        ("Search Keyword", 28),
        ("Source", 18),
        ("Notes", 60),
    ]

    for idx, (header, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx, value=header)
        cell.fill = EXCEL_HEADER_FILL
        cell.font = EXCEL_HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(idx)].width = width

    for row_idx, lead in enumerate(leads, start=2):
        values = [
            lead.company_name,
            lead.phone,
            lead.email,
            lead.website,
            lead.facebook_page,
            lead.city,
            lead.state,
            lead.search_keyword,
            lead.source,
            lead.notes,
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = EXCEL_BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if lead.website:
            ws.cell(row=row_idx, column=4).hyperlink = lead.website
            ws.cell(row=row_idx, column=4).style = "Hyperlink"
        if lead.facebook_page:
            ws.cell(row=row_idx, column=5).hyperlink = lead.facebook_page
            ws.cell(row=row_idx, column=5).style = "Hyperlink"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    summary = wb.create_sheet("Summary")
    summary_headers = [("Metric", 28), ("Value", 18)]
    for idx, (header, width) in enumerate(summary_headers, start=1):
        cell = summary.cell(row=1, column=idx, value=header)
        cell.fill = EXCEL_HEADER_FILL
        cell.font = EXCEL_HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        summary.column_dimensions[get_column_letter(idx)].width = width

    with_phone = sum(1 for lead in leads if lead.phone)
    with_email = sum(1 for lead in leads if lead.email)
    with_website = sum(1 for lead in leads if lead.website)
    metrics = [
        ("Exported At", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Total Leads", len(leads)),
        ("With Phone", with_phone),
        ("With Email", with_email),
        ("With Website", with_website),
        ("Keywords Used", ", ".join(normalize_search_terms(keywords))),
        ("Locations Used", ", ".join(f"{city}, {state}" for city, state in dedupe_location_pairs(locations))),
    ]
    for row_idx, (metric, value) in enumerate(metrics, start=2):
        summary.cell(row=row_idx, column=1, value=metric).font = EXCEL_BODY_FONT
        summary.cell(row=row_idx, column=2, value=value).font = EXCEL_BODY_FONT

    row = len(metrics) + 4
    summary.cell(row=row, column=1, value="City Summary").fill = EXCEL_HEADER_FILL
    summary.cell(row=row, column=1).font = EXCEL_HEADER_FONT
    summary.cell(row=row, column=2, value="Lead Count").fill = EXCEL_HEADER_FILL
    summary.cell(row=row, column=2).font = EXCEL_HEADER_FONT
    row += 1
    city_counts: dict[tuple[str, str], int] = defaultdict(int)
    for lead in leads:
        city_counts[(lead.city, lead.state)] += 1
    for city, state in dedupe_location_pairs(city_counts.keys()):
        summary.cell(row=row, column=1, value=f"{city}, {state}").font = EXCEL_BODY_FONT
        summary.cell(row=row, column=2, value=city_counts[(city, state)]).font = EXCEL_BODY_FONT
        row += 1

    row += 1
    summary.cell(row=row, column=1, value="Keyword Summary").fill = EXCEL_HEADER_FILL
    summary.cell(row=row, column=1).font = EXCEL_HEADER_FONT
    summary.cell(row=row, column=2, value="Lead Count").fill = EXCEL_HEADER_FILL
    summary.cell(row=row, column=2).font = EXCEL_HEADER_FONT
    row += 1
    keyword_counts: dict[str, int] = defaultdict(int)
    for lead in leads:
        keyword_counts[lead.search_keyword or "(unspecified)"] += 1
    for keyword in sorted(keyword_counts, key=lambda item: (-keyword_counts[item], item.lower())):
        summary.cell(row=row, column=1, value=keyword).font = EXCEL_BODY_FONT
        summary.cell(row=row, column=2, value=keyword_counts[keyword]).font = EXCEL_BODY_FONT
        row += 1

    settings = wb.create_sheet("Settings")
    settings_headers = [("Type", 18), ("Value", 42)]
    for idx, (header, width) in enumerate(settings_headers, start=1):
        cell = settings.cell(row=1, column=idx, value=header)
        cell.fill = EXCEL_HEADER_FILL
        cell.font = EXCEL_HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        settings.column_dimensions[get_column_letter(idx)].width = width

    row = 2
    for keyword in normalize_search_terms(keywords):
        settings.cell(row=row, column=1, value="Keyword").font = EXCEL_BODY_FONT
        settings.cell(row=row, column=2, value=keyword).font = EXCEL_BODY_FONT
        row += 1
    for city, state in dedupe_location_pairs(locations):
        settings.cell(row=row, column=1, value="Location").font = EXCEL_BODY_FONT
        settings.cell(row=row, column=2, value=f"{city}, {state}").font = EXCEL_BODY_FONT
        row += 1

    wb.save(filepath)
    return filepath


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def build_cli_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Facebook Lead Scraper")
    parser.add_argument("--search", default="construction companies", help="Single search term")
    parser.add_argument("--keywords", default="", help="Multiple keywords separated by semicolons or new lines")
    parser.add_argument("--city", default="", help="Single city override")
    parser.add_argument("--state", default="", help="State for --city")
    parser.add_argument(
        "--location",
        action="append",
        default=[],
        help="Repeatable city/state pair, e.g. --location 'Rockville,MD'",
    )
    parser.add_argument(
        "--locations",
        default="",
        help="Semicolon-separated city/state pairs, e.g. 'Rockville,MD;Gaithersburg,MD'",
    )
    parser.add_argument("--headless", action="store_true", help="Run browser in headless mode")
    parser.add_argument(
        "--results-per-city", "--max-results",
        dest="results_per_city",
        type=int,
        default=DEFAULT_RESULTS_PER_CITY,
        help="Target final leads per city (default: 10)",
    )
    parser.add_argument(
        "--candidate-depth",
        type=int,
        default=DEFAULT_CANDIDATE_SEARCH_DEPTH,
        help="How many candidate pages to inspect per city (default: 50)",
    )
    parser.add_argument("--user-data-dir", default=FACEBOOK_USER_DATA_DIR, help="Chrome user-data-dir to reuse an existing login session")
    parser.add_argument("--profile-directory", default=FACEBOOK_PROFILE_DIRECTORY, help="Chrome profile directory name, e.g. 'Default' or 'Profile 1'")
    parser.add_argument("--browser-binary", default=CHROME_BINARY, help="Explicit Chrome binary path if Selenium cannot detect Chrome")
    parser.add_argument("--debug-dir", default="", help="Directory to save HTML/screenshots for debugging")
    parser.add_argument("--output-dir", default="", help="Directory for timestamped Excel export")
    parser.add_argument("--export", action="store_true", help="Export results to a timestamped Excel file when done")
    return parser


def cli_main(argv: list[str] | None = None) -> int:
    parser = build_cli_parser()
    args = parser.parse_args(argv)

    if bool(_normalize_space(args.city)) ^ bool(_normalize_space(args.state)):
        parser.error("--city and --state must be used together")

    locations = build_locations_from_args(args)
    keywords = parse_keywords_text(args.keywords) or normalize_search_terms(args.search)

    log(f"Keywords ({len(keywords)}): " + "; ".join(keywords))
    log(f"Headless:    {args.headless}")
    log(f"FB Email:    {'SET' if FB_EMAIL else 'NOT SET'}")
    log(f"FB Password: {'SET' if FB_PASSWORD else 'NOT SET'}")
    log(f"Results/city:{args.results_per_city}")
    log(f"Candidate depth/city: {args.candidate_depth}")
    log(f"Locations ({len(locations)}): " + "; ".join(f"{city}, {state}" for city, state in locations))
    if args.user_data_dir:
        log(f"Chrome user-data-dir: {args.user_data_dir}")
    if args.profile_directory:
        log(f"Chrome profile:      {args.profile_directory}")
    if args.debug_dir:
        log(f"Debug dir:           {args.debug_dir}")
    print()

    leads = scrape_facebook_pages(
        search_terms=keywords,
        locations=locations,
        headless=args.headless,
        results_per_location=args.results_per_city,
        candidate_search_depth=args.candidate_depth,
        user_data_dir=args.user_data_dir,
        profile_directory=args.profile_directory,
        browser_binary=args.browser_binary,
        debug_dir=args.debug_dir,
    )
    print_results(leads)

    if args.export and leads:
        output_path = build_timestamped_excel_path(args.output_dir)
        export_leads_to_excel(leads, keywords, locations, output_path)
        log(f"Excel exported to: {output_path}")

    return 0


# ---------------------------------------------------------------------------
# GUI
# ---------------------------------------------------------------------------


class GuiConfig:
    def __init__(
        self,
        keywords: list[str],
        locations: list[tuple[str, str]],
        results_per_city: int,
        candidate_depth: int,
        headless: bool,
        user_data_dir: str,
        profile_directory: str,
        browser_binary: str,
        debug_dir: str,
        output_dir: str,
        auto_export: bool,
    ):
        self.keywords = keywords
        self.locations = locations
        self.results_per_city = results_per_city
        self.candidate_depth = candidate_depth
        self.headless = headless
        self.user_data_dir = user_data_dir
        self.profile_directory = profile_directory
        self.browser_binary = browser_binary
        self.debug_dir = debug_dir
        self.output_dir = output_dir
        self.auto_export = auto_export


def launch_gui():
    if tk is None or ttk is None:
        raise RuntimeError("Tkinter is not available in this Python environment.")

    class FacebookLeadScraperApp:
        def __init__(self, root: tk.Tk):
            self.root = root
            self.root.title("Facebook Lead Scraper")
            self.root.geometry("1440x920")
            self.root.minsize(1200, 760)

            self.message_queue: queue.Queue = queue.Queue()
            self.worker_thread: threading.Thread | None = None
            self.stop_event = threading.Event()
            self.current_leads: list[Lead] = []
            self.last_keywords: list[str] = list(DEFAULT_KEYWORDS)
            self.last_locations: list[tuple[str, str]] = list(DEFAULT_LOCATIONS)
            self._log_callback = lambda line: self.message_queue.put(("log", line))

            self.status_var = tk.StringVar(value="Ready")
            self.results_per_city_var = tk.IntVar(value=DEFAULT_RESULTS_PER_CITY)
            self.candidate_depth_var = tk.IntVar(value=DEFAULT_CANDIDATE_SEARCH_DEPTH)
            self.headless_var = tk.BooleanVar(value=False)
            self.auto_export_var = tk.BooleanVar(value=True)
            self.user_data_dir_var = tk.StringVar(value=FACEBOOK_USER_DATA_DIR)
            self.profile_directory_var = tk.StringVar(value=FACEBOOK_PROFILE_DIRECTORY or "Default")
            self.browser_binary_var = tk.StringVar(value=CHROME_BINARY)
            self.debug_dir_var = tk.StringVar(value="")
            self.output_dir_var = tk.StringVar(value=str(Path.cwd()))
            self.city_var = tk.StringVar()
            self.state_var = tk.StringVar(value="MD")

            self._build_ui()
            self._load_default_keywords()
            self._load_default_locations()
            self.root.after(150, self._process_queue)

        def _build_ui(self):
            self.root.columnconfigure(0, weight=1)
            self.root.rowconfigure(1, weight=1)

            toolbar = ttk.Frame(self.root, padding=(12, 10))
            toolbar.grid(row=0, column=0, sticky="ew")
            toolbar.columnconfigure(9, weight=1)

            ttk.Button(toolbar, text="Start", command=self.start_scrape).grid(row=0, column=0, padx=(0, 6))
            ttk.Button(toolbar, text="Stop", command=self.stop_scrape).grid(row=0, column=1, padx=(0, 6))
            ttk.Button(toolbar, text="Export Excel", command=self.export_results).grid(row=0, column=2, padx=(0, 12))
            ttk.Checkbutton(toolbar, text="Auto-export after run", variable=self.auto_export_var).grid(row=0, column=3, padx=(0, 12))
            ttk.Checkbutton(toolbar, text="Headless browser", variable=self.headless_var).grid(row=0, column=4, padx=(0, 12))
            ttk.Label(toolbar, text="Results / city:").grid(row=0, column=5, sticky="e")
            ttk.Spinbox(toolbar, from_=1, to=100, textvariable=self.results_per_city_var, width=6).grid(row=0, column=6, padx=(6, 12))
            ttk.Label(toolbar, text="Candidate depth:").grid(row=0, column=7, sticky="e")
            ttk.Spinbox(toolbar, from_=5, to=500, increment=5, textvariable=self.candidate_depth_var, width=7).grid(row=0, column=8, padx=(6, 0))

            notebook = ttk.Notebook(self.root)
            notebook.grid(row=1, column=0, sticky="nsew", padx=12, pady=(0, 8))

            config_tab = ttk.Frame(notebook, padding=12)
            results_tab = ttk.Frame(notebook, padding=12)
            logs_tab = ttk.Frame(notebook, padding=12)
            notebook.add(config_tab, text="Configuration")
            notebook.add(results_tab, text="Results")
            notebook.add(logs_tab, text="Logs")

            self._build_config_tab(config_tab)
            self._build_results_tab(results_tab)
            self._build_logs_tab(logs_tab)

            status_bar = ttk.Label(self.root, textvariable=self.status_var, anchor="w", relief="sunken", padding=(10, 6))
            status_bar.grid(row=2, column=0, sticky="ew")

        def _build_config_tab(self, parent):
            parent.columnconfigure(0, weight=1)
            parent.columnconfigure(1, weight=1)
            parent.rowconfigure(1, weight=1)

            keywords_frame = ttk.LabelFrame(parent, text="Keywords", padding=10)
            keywords_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 8), pady=(0, 8))
            keywords_frame.columnconfigure(0, weight=1)
            keywords_frame.rowconfigure(1, weight=1)
            ttk.Label(keywords_frame, text="Enter one keyword per line.").grid(row=0, column=0, sticky="w", pady=(0, 6))
            self.keywords_text = tk.Text(keywords_frame, height=10, wrap="word")
            self.keywords_text.grid(row=1, column=0, sticky="nsew")

            locations_frame = ttk.LabelFrame(parent, text="Cities", padding=10)
            locations_frame.grid(row=0, column=1, sticky="nsew", pady=(0, 8))
            locations_frame.columnconfigure(0, weight=1)
            locations_frame.rowconfigure(1, weight=1)

            entry_row = ttk.Frame(locations_frame)
            entry_row.grid(row=0, column=0, sticky="ew", pady=(0, 8))
            ttk.Label(entry_row, text="City").grid(row=0, column=0, sticky="w")
            ttk.Entry(entry_row, textvariable=self.city_var, width=24).grid(row=1, column=0, padx=(0, 8))
            ttk.Label(entry_row, text="State").grid(row=0, column=1, sticky="w")
            ttk.Entry(entry_row, textvariable=self.state_var, width=8).grid(row=1, column=1, padx=(0, 8))
            ttk.Button(entry_row, text="Add / Update", command=self.add_or_update_city).grid(row=1, column=2, padx=(0, 6))
            ttk.Button(entry_row, text="Remove Selected", command=self.remove_selected_cities).grid(row=1, column=3, padx=(0, 6))
            ttk.Button(entry_row, text="Load Defaults", command=self._load_default_locations).grid(row=1, column=4)

            tree_frame = ttk.Frame(locations_frame)
            tree_frame.grid(row=1, column=0, sticky="nsew")
            tree_frame.columnconfigure(0, weight=1)
            tree_frame.rowconfigure(0, weight=1)
            self.locations_tree = ttk.Treeview(tree_frame, columns=("city", "state"), show="headings", height=12, selectmode="extended")
            self.locations_tree.heading("city", text="City")
            self.locations_tree.heading("state", text="State")
            self.locations_tree.column("city", width=220, anchor="w")
            self.locations_tree.column("state", width=90, anchor="center")
            self.locations_tree.grid(row=0, column=0, sticky="nsew")
            loc_scroll = ttk.Scrollbar(tree_frame, orient="vertical", command=self.locations_tree.yview)
            loc_scroll.grid(row=0, column=1, sticky="ns")
            self.locations_tree.configure(yscrollcommand=loc_scroll.set)
            self.locations_tree.bind("<<TreeviewSelect>>", self.on_location_select)

            options_frame = ttk.LabelFrame(parent, text="Browser and output", padding=10)
            options_frame.grid(row=1, column=0, columnspan=2, sticky="nsew")
            options_frame.columnconfigure(1, weight=1)

            rows = [
                ("Chrome user-data-dir", self.user_data_dir_var, self.browse_user_data_dir),
                ("Chrome profile directory", self.profile_directory_var, None),
                ("Chrome binary", self.browser_binary_var, self.browse_browser_binary),
                ("Debug output directory", self.debug_dir_var, self.browse_debug_dir),
                ("Excel output directory", self.output_dir_var, self.browse_output_dir),
            ]
            for row_idx, (label, variable, browse_cmd) in enumerate(rows):
                ttk.Label(options_frame, text=label).grid(row=row_idx, column=0, sticky="w", pady=4, padx=(0, 8))
                ttk.Entry(options_frame, textvariable=variable).grid(row=row_idx, column=1, columnspan=3, sticky="ew", pady=4)
                if browse_cmd is not None:
                    ttk.Button(options_frame, text="Browse", command=browse_cmd).grid(row=row_idx, column=4, padx=(8, 0), pady=4)

        def _build_results_tab(self, parent):
            parent.columnconfigure(0, weight=1)
            parent.rowconfigure(0, weight=1)
            columns = ("company", "phone", "email", "website", "facebook", "city", "state", "keyword")
            self.results_tree = ttk.Treeview(parent, columns=columns, show="headings")
            headings = {
                "company": ("Company", 240),
                "phone": ("Phone", 120),
                "email": ("Email", 220),
                "website": ("Website", 220),
                "facebook": ("Facebook Page", 240),
                "city": ("City", 140),
                "state": ("State", 70),
                "keyword": ("Keyword", 180),
            }
            for key, (label, width) in headings.items():
                self.results_tree.heading(key, text=label)
                self.results_tree.column(key, width=width, anchor="w")
            self.results_tree.grid(row=0, column=0, sticky="nsew")
            y_scroll = ttk.Scrollbar(parent, orient="vertical", command=self.results_tree.yview)
            y_scroll.grid(row=0, column=1, sticky="ns")
            x_scroll = ttk.Scrollbar(parent, orient="horizontal", command=self.results_tree.xview)
            x_scroll.grid(row=1, column=0, sticky="ew")
            self.results_tree.configure(yscrollcommand=y_scroll.set, xscrollcommand=x_scroll.set)

        def _build_logs_tab(self, parent):
            parent.columnconfigure(0, weight=1)
            parent.rowconfigure(0, weight=1)
            self.logs_text = tk.Text(parent, wrap="word", state="disabled")
            self.logs_text.grid(row=0, column=0, sticky="nsew")
            scroll = ttk.Scrollbar(parent, orient="vertical", command=self.logs_text.yview)
            scroll.grid(row=0, column=1, sticky="ns")
            self.logs_text.configure(yscrollcommand=scroll.set)

        def _append_log(self, line: str):
            self.logs_text.configure(state="normal")
            self.logs_text.insert("end", line + "\n")
            self.logs_text.see("end")
            self.logs_text.configure(state="disabled")
            self.status_var.set(line)

        def _load_default_keywords(self):
            self.keywords_text.delete("1.0", "end")
            self.keywords_text.insert("1.0", "\n".join(DEFAULT_KEYWORDS))

        def _clear_locations(self):
            for item in self.locations_tree.get_children():
                self.locations_tree.delete(item)

        def _load_default_locations(self):
            self._clear_locations()
            for city, state in DEFAULT_LOCATIONS:
                self.locations_tree.insert("", "end", values=(city, state))

        def on_location_select(self, _event=None):
            selected = self.locations_tree.selection()
            if len(selected) != 1:
                return
            values = self.locations_tree.item(selected[0], "values")
            if len(values) >= 2:
                self.city_var.set(values[0])
                self.state_var.set(values[1])

        def add_or_update_city(self):
            city = _normalize_space(self.city_var.get())
            state = _normalize_space(self.state_var.get()).upper()
            if not city or not state:
                messagebox.showwarning("Missing city/state", "Enter both a city and state.")
                return

            selected = self.locations_tree.selection()
            if len(selected) == 1:
                self.locations_tree.item(selected[0], values=(city, state))
            else:
                self.locations_tree.insert("", "end", values=(city, state))
            self._dedupe_tree_locations()
            self.city_var.set("")

        def _dedupe_tree_locations(self):
            rows = []
            for item in self.locations_tree.get_children():
                city, state = self.locations_tree.item(item, "values")
                rows.append((city, state))
            rows = dedupe_location_pairs(rows)
            self._clear_locations()
            for city, state in rows:
                self.locations_tree.insert("", "end", values=(city, state))

        def remove_selected_cities(self):
            for item in self.locations_tree.selection():
                self.locations_tree.delete(item)

        def browse_user_data_dir(self):
            path = filedialog.askdirectory(title="Select Chrome user-data-dir")
            if path:
                self.user_data_dir_var.set(path)

        def browse_browser_binary(self):
            path = filedialog.askopenfilename(title="Select Chrome executable")
            if path:
                self.browser_binary_var.set(path)

        def browse_debug_dir(self):
            path = filedialog.askdirectory(title="Select debug output directory")
            if path:
                self.debug_dir_var.set(path)

        def browse_output_dir(self):
            path = filedialog.askdirectory(title="Select Excel output directory")
            if path:
                self.output_dir_var.set(path)

        def _gather_locations(self) -> list[tuple[str, str]]:
            rows = []
            for item in self.locations_tree.get_children():
                values = self.locations_tree.item(item, "values")
                if len(values) >= 2:
                    rows.append((values[0], values[1]))
            return dedupe_location_pairs(rows)

        def _gather_config(self) -> GuiConfig | None:
            keywords = parse_keywords_text(self.keywords_text.get("1.0", "end"))
            locations = self._gather_locations()

            if not keywords:
                messagebox.showwarning("Missing keywords", "Add at least one keyword.")
                return None
            if not locations:
                messagebox.showwarning("Missing cities", "Add at least one city/state pair.")
                return None

            try:
                results_per_city = max(1, int(self.results_per_city_var.get()))
                candidate_depth = max(results_per_city * 2, int(self.candidate_depth_var.get()))
            except Exception:
                messagebox.showwarning("Invalid numeric values", "Results per city and candidate depth must be numbers.")
                return None

            return GuiConfig(
                keywords=keywords,
                locations=locations,
                results_per_city=results_per_city,
                candidate_depth=candidate_depth,
                headless=bool(self.headless_var.get()),
                user_data_dir=_normalize_space(self.user_data_dir_var.get()),
                profile_directory=_normalize_space(self.profile_directory_var.get()),
                browser_binary=_normalize_space(self.browser_binary_var.get()),
                debug_dir=_normalize_space(self.debug_dir_var.get()),
                output_dir=_normalize_space(self.output_dir_var.get()),
                auto_export=bool(self.auto_export_var.get()),
            )

        def _set_running(self, running: bool):
            self.status_var.set("Running..." if running else "Ready")

        def start_scrape(self):
            if self.worker_thread and self.worker_thread.is_alive():
                messagebox.showinfo("Already running", "A scrape is already in progress.")
                return

            config = self._gather_config()
            if config is None:
                return

            self.current_leads = []
            for item in self.results_tree.get_children():
                self.results_tree.delete(item)
            self.logs_text.configure(state="normal")
            self.logs_text.delete("1.0", "end")
            self.logs_text.configure(state="disabled")

            self.last_keywords = list(config.keywords)
            self.last_locations = list(config.locations)
            self.stop_event.clear()
            self._set_running(True)

            self.worker_thread = threading.Thread(target=self._worker_run, args=(config,), daemon=True)
            self.worker_thread.start()

        def stop_scrape(self):
            if self.worker_thread and self.worker_thread.is_alive():
                self.stop_event.set()
                self.status_var.set("Stop requested...")

        def _worker_run(self, config: GuiConfig):
            add_log_subscriber(self._log_callback)
            try:
                leads = scrape_facebook_pages(
                    search_terms=config.keywords,
                    locations=config.locations,
                    headless=config.headless,
                    results_per_location=config.results_per_city,
                    candidate_search_depth=config.candidate_depth,
                    user_data_dir=config.user_data_dir,
                    profile_directory=config.profile_directory,
                    browser_binary=config.browser_binary,
                    debug_dir=config.debug_dir,
                    stop_event=self.stop_event,
                )
                exported_path = ""
                if config.auto_export and leads:
                    output_path = build_timestamped_excel_path(config.output_dir)
                    export_leads_to_excel(leads, config.keywords, config.locations, output_path)
                    exported_path = str(output_path)
                self.message_queue.put(("done", {
                    "leads": leads,
                    "keywords": config.keywords,
                    "locations": config.locations,
                    "exported_path": exported_path,
                    "stopped": self.stop_event.is_set(),
                }))
            except Exception as exc:
                self.message_queue.put(("error", str(exc)))
            finally:
                remove_log_subscriber(self._log_callback)

        def _process_queue(self):
            try:
                while True:
                    kind, payload = self.message_queue.get_nowait()
                    if kind == "log":
                        self._append_log(payload)
                    elif kind == "done":
                        self.current_leads = payload["leads"]
                        self.last_keywords = payload["keywords"]
                        self.last_locations = payload["locations"]
                        self._populate_results(self.current_leads)
                        self._set_running(False)
                        if payload.get("exported_path"):
                            self.status_var.set(f"Completed and exported to {payload['exported_path']}")
                            messagebox.showinfo("Export complete", f"Excel file created:\n{payload['exported_path']}")
                        elif payload.get("stopped"):
                            self.status_var.set("Stopped")
                        else:
                            self.status_var.set(f"Completed — {len(self.current_leads)} lead(s)")
                    elif kind == "error":
                        self._set_running(False)
                        self.status_var.set("Error")
                        messagebox.showerror("Scrape error", payload)
            except queue.Empty:
                pass
            finally:
                self.root.after(150, self._process_queue)

        def _populate_results(self, leads: list[Lead]):
            for item in self.results_tree.get_children():
                self.results_tree.delete(item)
            for lead in leads:
                self.results_tree.insert(
                    "",
                    "end",
                    values=(
                        lead.company_name,
                        lead.phone,
                        lead.email,
                        lead.website,
                        lead.facebook_page,
                        lead.city,
                        lead.state,
                        lead.search_keyword,
                    ),
                )

        def export_results(self):
            if not self.current_leads:
                messagebox.showinfo("No results", "Run a scrape first, then export the results.")
                return
            output_path = build_timestamped_excel_path(self.output_dir_var.get())
            export_leads_to_excel(self.current_leads, self.last_keywords, self.last_locations, output_path)
            self.status_var.set(f"Excel exported to {output_path}")
            messagebox.showinfo("Export complete", f"Excel file created:\n{output_path}")

    root = tk.Tk()
    try:
        style = ttk.Style(root)
        with suppress(Exception):
            style.theme_use("clam")
    except Exception:
        pass
    FacebookLeadScraperApp(root)
    root.mainloop()


def main() -> int:
    # Keep execution focused on the Facebook scraping module (CLI only).
    return cli_main()


if __name__ == "__main__":
    raise SystemExit(main())

