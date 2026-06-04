import argparse
import sys
import types
import unittest
from pathlib import Path
from unittest.mock import patch


def _install_openpyxl_stub_if_missing() -> None:
    try:
        import openpyxl  # noqa: F401
        return
    except Exception:
        pass

    openpyxl_module = types.ModuleType("openpyxl")
    styles_module = types.ModuleType("openpyxl.styles")
    utils_module = types.ModuleType("openpyxl.utils")

    class _DummyWorkbook:
        pass

    class _DummyStyle:
        def __init__(self, *args, **kwargs):
            self.args = args
            self.kwargs = kwargs

    def _dummy_get_column_letter(index: int) -> str:
        return str(index)

    openpyxl_module.Workbook = _DummyWorkbook
    styles_module.Alignment = _DummyStyle
    styles_module.Font = _DummyStyle
    styles_module.PatternFill = _DummyStyle
    utils_module.get_column_letter = _dummy_get_column_letter

    sys.modules["openpyxl"] = openpyxl_module
    sys.modules["openpyxl.styles"] = styles_module
    sys.modules["openpyxl.utils"] = utils_module


_install_openpyxl_stub_if_missing()

import testing


class TestFacebookCliHelpers(unittest.TestCase):
    def test_parse_keywords_text_dedupes_and_trims(self):
        raw = " Roofers ;\nroofers\n General Contractors ;  "
        actual = testing.parse_keywords_text(raw)
        self.assertEqual(actual, ["Roofers", "General Contractors"])

    def test_build_locations_from_args_uses_defaults_when_empty(self):
        args = argparse.Namespace(location=[], locations="", city="", state="")
        actual = testing.build_locations_from_args(args)
        self.assertEqual(actual, list(testing.DEFAULT_LOCATIONS))

    def test_build_locations_from_args_combines_and_dedupes(self):
        args = argparse.Namespace(
            location=["Rockville,MD", "Arlington|VA"],
            locations="Rockville,MD;Bethesda,MD",
            city="Arlington",
            state="VA",
        )
        actual = testing.build_locations_from_args(args)
        self.assertEqual(
            actual,
            [("Rockville", "MD"), ("Arlington", "VA"), ("Bethesda", "MD")],
        )


class TestFacebookCliMain(unittest.TestCase):
    def test_cli_main_calls_scraper_and_exports(self):
        fake_leads = [
            testing.Lead(
                company_name="Acme Construction",
                city="Rockville",
                state="MD",
                search_keyword="roofers",
                facebook_page="https://www.facebook.com/acme",
            )
        ]
        output_file = Path("C:/temp/fb_leads_test.xlsx")

        with patch("testing.scrape_facebook_pages", return_value=fake_leads) as mock_scrape, patch(
            "testing.print_results"
        ) as mock_print_results, patch(
            "testing.build_timestamped_excel_path", return_value=output_file
        ) as mock_output_path, patch(
            "testing.export_leads_to_excel"
        ) as mock_export:
            rc = testing.cli_main(
                [
                    "--keywords",
                    "roofers; painters",
                    "--location",
                    "Rockville,MD",
                    "--results-per-city",
                    "3",
                    "--candidate-depth",
                    "8",
                    "--headless",
                    "--export",
                ]
            )

        self.assertEqual(rc, 0)
        mock_scrape.assert_called_once()
        mock_print_results.assert_called_once_with(fake_leads)
        mock_output_path.assert_called_once_with("")
        mock_export.assert_called_once_with(
            fake_leads,
            ["roofers", "painters"],
            [("Rockville", "MD")],
            output_file,
        )

    def test_cli_main_requires_city_and_state_together(self):
        with self.assertRaises(SystemExit):
            testing.cli_main(["--city", "Rockville"])


if __name__ == "__main__":
    unittest.main()
